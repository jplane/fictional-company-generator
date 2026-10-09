import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT = Path(__file__).resolve().parents[2]
plan = json.loads((ROOT / "build" / "plan.json").read_text())
selected = set(sys.argv[1:])
USD = '$#,##0;($#,##0);"-"'
NUM = '#,##0;(#,##0);"-"'
PCT = '0.0%;(0.0%);"-"'
DARK, COPPER, CREAM = "18242D", "B5683A", "F3F0E9"
FIELDS = [
    ("revenue_usd", "Recognized revenue"), ("cogs_usd", "Cost of goods sold"),
    ("opex_usd", "Operating expenses"), ("operating_income_usd", "Operating result"),
    ("headcount", "Period-end headcount"), ("engineering_services_usd", "Engineering services"),
    ("instrument_development_usd", "Instrument development"), ("mission_analysis_usd", "Mission analysis"),
    ("research_development_usd", "Research and development"), ("program_operations_usd", "Program operations"),
    ("corporate_administration_usd", "Corporate administration"), ("contracted_backlog_usd", "Contracted backlog"),
    ("completed_work_packages", "Completed work packages"), ("accepted_reference_runs", "Accepted reference runs"),
    ("rejected_reference_runs", "Rejected reference runs"),
]
ROWS = {key: n + 8 for n, (key, label) in enumerate(FIELDS)}


def format_sheet(ws, title, date, units="USD unless stated"):
    ws.sheet_view.showGridLines = False
    ws.merge_cells("A1:H2")
    ws["A1"] = title
    ws["A1"].font = Font(name="Arial", size=17, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor=DARK)
    ws["A1"].alignment = Alignment(wrap_text=True, vertical="center")
    ws["A3"] = f"Published {date} | {units} | Internal"
    ws["A3"].font = Font(name="Arial", size=10, color="53616A")
    ws.freeze_panes = "B8"
    ws.column_dimensions["A"].width = 38
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["C"].width = 22
    for col in "DEFGH":
        ws.column_dimensions[col].width = 19
    ws.row_dimensions[1].height = 24
    ws.row_dimensions[2].height = 24
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.print_options.horizontalCentered = True
    ws.print_area = "A1:H36"
    ws.oddFooter.center.text = "AstroDig LLC | Internal | &P"


def text_block(ws, row, text, height=46):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=8)
    cell = ws.cell(row, 1, text)
    cell.alignment = Alignment(wrap_text=True, vertical="top")
    cell.font = Font(name="Arial", size=11, color=DARK)
    ws.row_dimensions[row].height = height


def formula(ws, cell, expression, fmt=USD):
    c = ws[cell]
    c.value = expression
    c.number_format = fmt
    c.font = Font(name="Arial", size=11, color="008000" if "!" in expression else "000000")


def source_ref(key, col="B"):
    return f"'Source actuals'!{col}{ROWS[key]}"


def performance_module(ws, number):
    if number == 1:
        for row, key in enumerate(["engineering_services_usd", "instrument_development_usd", "mission_analysis_usd"], 8):
            ws.cell(row, 1, dict(FIELDS)[key])
            formula(ws, f"B{row}", "=" + source_ref(key))
            formula(ws, f"C{row}", f"=IF(B11=0,0,B{row}/B11)", PCT)
        ws["A11"], ws["A12"] = "Service total", "Recognized revenue"
        formula(ws, "B11", "=SUM(B8:B10)")
        formula(ws, "B12", "=" + source_ref("revenue_usd"))
        ws["A14"] = "Service reconciliation"
        formula(ws, "B14", "=B11-B12")
        ws["C14"] = '=IF(B14=0,"RECONCILED","REVIEW")'
        chart = BarChart()
        chart.title = "Contracted service revenue"
        chart.y_axis.scaling.min = 0
        chart.add_data(Reference(ws, min_col=2, min_row=7, max_row=10), titles_from_data=True)
        chart.set_categories(Reference(ws, min_col=1, min_row=8, max_row=10))
        chart.width, chart.height = 15, 8
        ws.add_chart(chart, "D8")
    elif number == 2:
        entries = [
            ("Research spending", "=" + source_ref("research_development_usd"), USD),
            ("Research / revenue", "=" + source_ref("research_development_usd") + "/" + source_ref("revenue_usd"), PCT),
            ("Research / period-end employee", "=" + source_ref("research_development_usd") + "/" + source_ref("headcount"), USD),
            ("Research / operating expenses", "=" + source_ref("research_development_usd") + "/" + source_ref("opex_usd"), PCT),
            ("Program operations", "=" + source_ref("program_operations_usd"), USD),
            ("Administration", "=" + source_ref("corporate_administration_usd"), USD),
            ("Expense allocation difference", "=B8+B12+B13-" + source_ref("opex_usd"), USD),
        ]
        for row, (label, expression, fmt) in enumerate(entries, 8):
            ws.cell(row, 1, label)
            formula(ws, f"B{row}", expression, fmt)
        ws["C14"] = '=IF(B14=0,"RECONCILED","REVIEW")'
    elif number == 3:
        entries = [
            ("Period-end headcount", "=" + source_ref("headcount"), NUM),
            ("Completed work packages", "=" + source_ref("completed_work_packages"), NUM),
            ("Work packages / employee", "=B9/B8", "0.00"),
            ("Accepted reference runs", "=" + source_ref("accepted_reference_runs"), NUM),
            ("Rejected reference runs", "=" + source_ref("rejected_reference_runs"), NUM),
            ("Recorded reference runs", "=SUM(B11:B12)", NUM),
            ("Reference acceptance fraction", "=B11/B13", PCT),
        ]
        for row, (label, expression, fmt) in enumerate(entries, 8):
            ws.cell(row, 1, label)
            formula(ws, f"B{row}", expression, fmt)
        ws["A16"] = "Run-count difference"
        formula(ws, "B16", "=B13-B11-B12", NUM)
        ws["C16"] = '=IF(B16=0,"RECONCILED","REVIEW")'
    elif number == 4:
        entries = [
            ("Q3 recognized revenue", "=" + source_ref("revenue_usd", "G"), USD),
            ("Q2 recognized revenue", "=" + source_ref("revenue_usd", "F"), USD),
            ("Revenue movement", "=B8-B9", USD),
            ("Revenue growth", "=B10/B9", PCT),
            ("Q3 operating expenses", "=" + source_ref("opex_usd", "G"), USD),
            ("Q2 operating expenses", "=" + source_ref("opex_usd", "F"), USD),
            ("Derived operating-result movement", "=B10-(" + source_ref("cogs_usd", "G") + "-" + source_ref("cogs_usd", "F") + ")-(B12-B13)", USD),
            ("Reported operating-result movement", "=" + source_ref("operating_income_usd", "G") + "-" + source_ref("operating_income_usd", "F"), USD),
        ]
        for row, (label, expression, fmt) in enumerate(entries, 8):
            ws.cell(row, 1, label)
            formula(ws, f"B{row}", expression, fmt)
        ws["A17"] = "Operating bridge difference"
        formula(ws, "B17", "=B14-B15")
        ws["C17"] = '=IF(B17=0,"RECONCILED","REVIEW")'
        chart = BarChart()
        chart.title = "Revenue recognition: Q2 and Q3"
        chart.y_axis.scaling.min = 0
        chart.add_data(Reference(ws, min_col=2, min_row=7, max_row=9), titles_from_data=True)
        chart.set_categories(Reference(ws, min_col=1, min_row=8, max_row=9))
        chart.width, chart.height = 15, 8
        ws.add_chart(chart, "D8")
    elif number == 5:
        entries = [
            ("Contracted backlog", "=" + source_ref("contracted_backlog_usd"), USD),
            ("Current-quarter revenue", "=" + source_ref("revenue_usd"), USD),
            ("Backlog / quarter revenue", "=B8/B9", '0.0"x"'),
            ("Period-end headcount", "=" + source_ref("headcount"), NUM),
            ("Backlog / period-end employee", "=B8/B11", USD),
        ]
        for row, (label, expression, fmt) in enumerate(entries, 8):
            ws.cell(row, 1, label)
            formula(ws, f"B{row}", expression, fmt)
        ws["A14"] = "Backlog source difference"
        formula(ws, "B14", "=B8-" + source_ref("contracted_backlog_usd"))
        ws["C14"] = '=IF(B14=0,"RECONCILED","REVIEW")'


def finalize(wb):
    for ws in wb:
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is None:
                    continue
                if cell.font.name != "Arial":
                    cell.font = Font(name="Arial", size=11, color="000000")
                if cell.row >= 7 and cell.column == 1 and not ws.merged_cells:
                    cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.row_dimensions[7].height = 23
        for cell in ws[7]:
            if cell.value is not None:
                cell.fill = PatternFill("solid", fgColor=COPPER)
                cell.font = Font(name="Arial", size=11, bold=True, color="FFFFFF")


def financial(item, draft):
    m = draft["model"]
    a = m["anchor"]
    wb = Workbook()
    ws = wb.active
    ws.title = "Source actuals"
    format_sheet(ws, draft["title"], draft["publication_date"])
    ws.column_dimensions["C"].width = 12
    ws.column_dimensions["D"].width = 33
    for col in "EFG":
        ws.column_dimensions[col].width = 22
    text_block(ws, 5, f'Q{a["quarter"]} {a["year"]}; period ended {a["period_end"]}. '
               "Company operating ledger and campaign register. Headcount and backlog are period-end snapshots.")
    for col, val in enumerate(["Measure", "Current period", "Units", "Available-quarter measure"], 1):
        ws.cell(7, col, val)
    for n, q in enumerate(m["quarters"], 5):
        ws.cell(7, n, f'Q{q["quarter"]} {q["year"]}')
    count_fields = {"headcount", "completed_work_packages", "accepted_reference_runs", "rejected_reference_runs"}
    for key, label in FIELDS:
        row = ROWS[key]
        ws.cell(row, 1, label)
        c = ws.cell(row, 2, a[key])
        c.font = Font(name="Arial", size=11, color="0000FF")
        c.number_format = NUM if key in count_fields else USD
        c.comment = Comment(f'Source: AstroDig operating ledger, {m["source_date"]}, '
                            f'Q{a["quarter"]} {a["year"]}, {label}.', "AstroDig Finance")
        ws.cell(row, 3, "count" if key in count_fields else "USD")
        ws.cell(row, 3).alignment = Alignment(indent=1)
        ws.cell(row, 4, label)
        for n, q in enumerate(m["quarters"], 5):
            c = ws.cell(row, n, q[key])
            c.font = Font(name="Arial", size=11, color="0000FF")
            c.number_format = NUM if key in count_fields else USD
            c.comment = Comment(f'Source: AstroDig operating ledger, Q{q["quarter"]} {q["year"]}, {label}.',
                                "AstroDig Finance")
    text_block(ws, 25, a["revenue_policy"], 40)
    text_block(ws, 28, "Revenue is recognized against accepted contracted work. No forecast, cash balance, "
               "resource valuation, or projected asteroid-sales input is included.", 40)
    for module_index, (name, question, interpretation) in enumerate(m["modules"]):
        performance = item["type"] == "financial_model"
        sheet_names = (["Income bridge", "Service economics", "Development intensity", "Delivery capacity", "Quarter bridge", "Backlog commitments"]
                       if performance else ["Income bridge", "Margins", "Revenue mix", "Cost drivers", "Productivity", "Trends and commitments"])
        sheet_name = sheet_names[module_index]
        s = wb.create_sheet(sheet_name)
        format_sheet(s, name, draft["publication_date"])
        text_block(s, 5, question, 32)
        s["A7"], s["B7"], s["C7"] = "Measure", "Value", "Check / share"
        if module_index == 0:
            for row, key in enumerate(["revenue_usd", "cogs_usd", "opex_usd"], 8):
                s.cell(row, 1, dict(FIELDS)[key])
                formula(s, f"B{row}", "=" + source_ref(key))
            s["A11"] = "Derived operating result"
            formula(s, "B11", "=B8-B9-B10")
            s["A12"] = "Reported operating result"
            formula(s, "B12", "=" + source_ref("operating_income_usd"))
            s["A14"] = "Reconciliation difference"
            formula(s, "B14", "=B11-B12")
            s["C14"] = '=IF(B14=0,"RECONCILED","REVIEW")'
        elif performance:
            performance_module(s, module_index)
        elif module_index == 1:
            entries = [
                ("Revenue", "=" + source_ref("revenue_usd"), USD),
                ("Gross profit", "=" + source_ref("revenue_usd") + "-" + source_ref("cogs_usd"), USD),
                ("Gross margin", "=IF(B8=0,0,B9/B8)", PCT),
                ("Operating margin", "=IF(B8=0,0," + source_ref("operating_income_usd") + "/B8)", PCT),
                ("Opex / revenue", "=IF(B8=0,0," + source_ref("opex_usd") + "/B8)", PCT),
                ("COGS / revenue", "=IF(B8=0,0," + source_ref("cogs_usd") + "/B8)", PCT),
            ]
            for row, (label, expression, fmt) in enumerate(entries, 8):
                s.cell(row, 1, label)
                formula(s, f"B{row}", expression, fmt)
            s["A15"] = "Gross-margin balance check"
            formula(s, "B15", "=B10+B13-1", PCT)
            s["C15"] = '=IF(ABS(B15)<0.000001,"RECONCILED","REVIEW")'
        elif module_index in (2, 3):
            keys = (["engineering_services_usd", "instrument_development_usd", "mission_analysis_usd"]
                    if module_index == 2 else ["research_development_usd", "program_operations_usd", "corporate_administration_usd"])
            for row, key in enumerate(keys, 8):
                s.cell(row, 1, dict(FIELDS)[key])
                formula(s, f"B{row}", "=" + source_ref(key))
                formula(s, f"C{row}", f"=IF(B11=0,0,B{row}/B11)", PCT)
            s["A11"] = "Allocated total"
            formula(s, "B11", "=SUM(B8:B10)")
            formula(s, "C11", "=SUM(C8:C10)", PCT)
            s["A12"] = "Reported total"
            formula(s, "B12", "=" + source_ref("revenue_usd" if module_index == 2 else "opex_usd"))
            s["A14"] = "Allocation difference"
            formula(s, "B14", "=B11-B12")
            s["C14"] = '=IF(B14=0,"RECONCILED","REVIEW")'
            chart = BarChart()
            chart.title = "Recognized revenue mix" if module_index == 2 else "Operating expense allocation"
            chart.y_axis.scaling.min = 0
            chart.add_data(Reference(s, min_col=2, min_row=7, max_row=10), titles_from_data=True)
            chart.set_categories(Reference(s, min_col=1, min_row=8, max_row=10))
            chart.width, chart.height = 15, 8
            s.add_chart(chart, "D8")
        elif module_index == 4:
            entries = [
                ("Period-end headcount", "=" + source_ref("headcount"), NUM),
                ("Revenue / period-end employee", "=" + source_ref("revenue_usd") + "/B8", USD),
                ("Completed work packages", "=" + source_ref("completed_work_packages"), NUM),
                ("Work packages / employee", "=B10/B8", "0.00"),
                ("Accepted reference runs", "=" + source_ref("accepted_reference_runs"), NUM),
                ("Rejected reference runs", "=" + source_ref("rejected_reference_runs"), NUM),
                ("Total recorded runs", "=SUM(B12:B13)", NUM),
                ("Reference acceptance fraction", "=IF(B14=0,0,B12/B14)", PCT),
            ]
            for row, (label, expression, fmt) in enumerate(entries, 8):
                s.cell(row, 1, label)
                formula(s, f"B{row}", expression, fmt)
            s["A17"] = "Run-count reconciliation"
            formula(s, "B17", "=B14-B12-B13", NUM)
            s["C17"] = '=IF(B17=0,"RECONCILED","REVIEW")'
        else:
            for col, q in enumerate(m["quarters"], 2):
                letter = get_column_letter(col)
                source_col = get_column_letter(col + 3)
                s.cell(7, col, f'Q{q["quarter"]} {q["year"]}')
                for row, key in enumerate(["revenue_usd", "cogs_usd", "opex_usd", "operating_income_usd", "contracted_backlog_usd", "headcount"], 8):
                    s.cell(row, 1, dict(FIELDS)[key])
                    formula(s, f"{letter}{row}", "=" + source_ref(key, source_col),
                            NUM if key == "headcount" else USD)
                s["A15"] = "Quarter income difference"
                formula(s, f"{letter}15", f"={letter}8-{letter}9-{letter}10-{letter}11")
                s.cell(16, col, f'=IF({letter}15=0,"RECONCILED","REVIEW")')
            s["A18"] = "Available-quarter revenue sum"
            formula(s, "B18", f"=SUM(B8:{get_column_letter(len(m['quarters'])+1)}8)")
            if len(m["quarters"]) > 1:
                s["A19"] = "Latest-quarter revenue growth"
                last = get_column_letter(len(m["quarters"]) + 1)
                prev = get_column_letter(len(m["quarters"]))
                formula(s, "B19", f"={last}8/{prev}8-1", PCT)
            else:
                s["A19"], s["B19"] = "Comparison availability", "Founding-quarter baseline only"
            s["A21"] = "Backlog interpretation"
            s["B21"] = "Remaining work; not cash"
        text_block(s, 26, interpretation, 60)
        text_block(s, 30, "Control: formulas link to source actuals; reconciliation checks must read RECONCILED. "
                   "Ratios describe the reported period and do not constitute forecast or mission readiness.", 42)
    notes = wb.create_sheet("Methodology")
    format_sheet(notes, "Methodology and evidence boundaries", draft["publication_date"])
    row = 5
    for section in draft["sections"]:
        text_block(notes, row, section["heading"] + ": " + section["text"], 62)
        row += 3
    notes.print_area = f"A1:H{row}"
    notes.page_setup.fitToHeight = 2
    finalize(wb)
    return wb


def specification(item, draft):
    wb = Workbook()
    ws = wb.active
    ws.title = "Requirements"
    format_sheet(ws, draft["title"], draft["publication_date"], "Units as specified")
    columns = ["Requirement", "Control", "Requirement statement", "Acceptance criterion",
               "Method", "Dependency", "Disposition", "Owner"]
    for c, value in enumerate(columns, 1):
        ws.cell(7, c, value)
    for row, req in enumerate(draft["requirement_rows"], 8):
        values = [req["requirement_id"], req["group"], req["statement"], req["acceptance_criterion"],
                  req["verification_method"], req["interface"], req["failure_disposition"], "Sasha Linden"]
        for c, value in enumerate(values, 1):
            ws.cell(row, c, value)
            ws.cell(row, c).alignment = Alignment(wrap_text=True, vertical="top", indent=1)
        ws.row_dimensions[row].height = 78
    for col, width in zip("ABCDEFGH", [15, 22, 38, 42, 22, 24, 38, 20]):
        ws.column_dimensions[col].width = width
    table = Table(displayName="AcceptanceRequirements", ref=f"A7:H{7+len(draft['requirement_rows'])}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(table)
    ws.auto_filter.ref = table.ref
    ws.print_area = f"A1:H{7+len(draft['requirement_rows'])}"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 4
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.print_title_rows = "7:7"
    notes = wb.create_sheet("Scope and release gate")
    format_sheet(notes, "Scope, dependencies, and acceptance", draft["publication_date"], "Units as specified")
    row = 5
    for section in draft["sections"]:
        if section["role"] == "requirements":
            continue
        text_block(notes, row, section["heading"] + ": " + section["text"], 92)
        row += 5
    finalize(wb)
    ws.row_dimensions[7].height = 32
    for cell in ws[7]:
        cell.alignment = Alignment(wrap_text=True, vertical="center", indent=1)
    return wb


for item in plan["documents"]:
    if item["format"] != "xlsx" or selected and item["id"] not in selected:
        continue
    output = ROOT / "documents" / item["filename"]
    if output.exists():
        continue
    draft = json.loads((ROOT / "build" / "drafts" / f'{item["id"]}.json').read_text())
    wb = specification(item, draft) if item["type"] == "product_spec" else financial(item, draft)
    wb.properties.creator = "AstroDig LLC"
    wb.properties.title = draft["title"]
    wb.properties.subject = "Controlled technical specification" if item["type"] == "product_spec" else "USD operating performance"
    wb.calculation.fullCalcOnLoad = True
    wb.save(output)
    print(item["id"], item["filename"])
