import importlib.util
import json
import os
import re
import subprocess
import sys
import zipfile
from pathlib import Path
from xml.etree import ElementTree

import pymupdf as fitz
from openpyxl import load_workbook
from PIL import Image, ImageDraw

ROOT = Path(__file__).resolve().parents[2]
BUILD = ROOT / "build"
SOFFICE = Path(r"C:\Program Files\LibreOffice\program\soffice.com")
POPPLER = Path(os.environ["LOCALAPPDATA"]) / "Microsoft" / "WinGet" / "Packages" / "oschwartz10612.Poppler_Microsoft.Winget.Source_8wekyb3d8bbwe" / "poppler-25.07.0" / "Library" / "bin" / "pdftoppm.exe"
spec = importlib.util.spec_from_file_location(
    "company", ROOT.parents[1] / ".github" / "skills" / "company-content-generator" / "scripts" / "company.py")
company = importlib.util.module_from_spec(spec)
spec.loader.exec_module(company)
plan = company.load(BUILD / "plan.json")
selected = set(sys.argv[1:])
items = [d for d in plan["documents"] if not selected or d["id"] in selected]
profile = (BUILD / "libreoffice-profile").as_uri()
renders = BUILD / "renders"
renders.mkdir(exist_ok=True)
recalculated = BUILD / "recalculated"
recalculated.mkdir(exist_ok=True)
checks_dir = BUILD / "checks"
checks_dir.mkdir(exist_ok=True)


def run_lo(args):
    result = subprocess.run([str(SOFFICE), f"-env:UserInstallation={profile}",
                             "--headless", "--norestore"] + args,
                            capture_output=True, text=True, timeout=900)
    (BUILD / "render.log").open("a", encoding="utf-8").write(result.stdout + result.stderr)
    if result.returncode:
        raise RuntimeError(result.stdout + result.stderr)


def normalize(text):
    text = re.sub(r"(?<=\w)-\s+(?=\w)", "-", text)
    return re.sub(r"\s+", " ", text).strip()


def check_office_text(item, draft, path):
    with zipfile.ZipFile(path) as archive:
        if item["format"] == "docx":
            tree = ElementTree.fromstring(archive.read("word/document.xml"))
            text = " ".join(n.text or "" for n in tree.iter() if n.tag.endswith("}t"))
            expected = [draft["title"]] + [s["text"] for s in draft["sections"]]
            for section in expected:
                assert normalize(section) in normalize(text), (item["id"], "Word text mismatch")
        elif item["format"] == "pptx":
            slides = sorted([n for n in archive.namelist() if re.fullmatch(r"ppt/slides/slide\d+\.xml", n)],
                            key=lambda n: int(re.search(r"slide(\d+)", n).group(1)))
            assert len(slides) == len(draft["slides"])
            for name, slide in zip(slides, draft["slides"]):
                tree = ElementTree.fromstring(archive.read(name))
                text = " ".join(n.text or "" for n in tree.iter() if n.tag.endswith("}t"))
                for body in slide["body"]:
                    assert normalize(body) in normalize(text), (item["id"], name, body)
        for name in archive.namelist():
            if name.endswith(".xml"):
                contents = archive.read(name).decode("utf-8")
                assert not re.search(r"\b(?:fictional|simulated|AI-generated|lorem ipsum)\b", contents, re.I), name
    company.audit_package(path, item["format"])


def check_workbook(item, draft, path):
    formulas = load_workbook(path, data_only=False)
    values = load_workbook(path, data_only=True)
    formula_count = 0
    for ws in formulas:
        for row in ws:
            for cell in row:
                cached = values[ws.title][cell.coordinate]
                assert cached.data_type != "e", (item["id"], ws.title, cell.coordinate, cached.value)
                if cell.data_type == "f":
                    formula_count += 1
                    assert cached.value is not None, (item["id"], "Missing cache", ws.title, cell.coordinate)
                    assert "[" not in cell.value, (item["id"], "External workbook link")
                    assert cell.value.startswith("=")
    if item["type"] != "product_spec":
        a = draft["model"]["anchor"]
        expected = {
            ("Income bridge", "B11"): a["operating_income_usd"],
            ("Income bridge", "B14"): 0,
            ("Margins", "B10"): (a["revenue_usd"] - a["cogs_usd"]) / a["revenue_usd"],
            ("Margins", "B11"): a["operating_income_usd"] / a["revenue_usd"],
            ("Revenue mix", "B11"): a["revenue_usd"],
            ("Revenue mix", "B14"): 0,
            ("Cost drivers", "B11"): a["opex_usd"],
            ("Cost drivers", "B14"): 0,
            ("Productivity", "B9"): a["revenue_usd"] / a["headcount"],
            ("Productivity", "B14"): a["accepted_reference_runs"] + a["rejected_reference_runs"],
            ("Productivity", "B15"): a["accepted_reference_runs"] / (a["accepted_reference_runs"] + a["rejected_reference_runs"]),
            ("Productivity", "B17"): 0,
            ("Trends and commitments", "B18"): sum(q["revenue_usd"] for q in draft["model"]["quarters"]),
        }
        if item["type"] == "financial_model":
            q2, q3 = draft["model"]["quarters"][1:3]
            expected = {
                ("Income bridge", "B11"): a["operating_income_usd"],
                ("Income bridge", "B14"): 0,
                ("Service economics", "B11"): a["revenue_usd"],
                ("Service economics", "B14"): 0,
                ("Development intensity", "B9"): a["research_development_usd"] / a["revenue_usd"],
                ("Development intensity", "B10"): a["research_development_usd"] / a["headcount"],
                ("Development intensity", "B14"): 0,
                ("Delivery capacity", "B14"): a["accepted_reference_runs"] / (a["accepted_reference_runs"] + a["rejected_reference_runs"]),
                ("Delivery capacity", "B16"): 0,
                ("Quarter bridge", "B14"): q3["operating_income_usd"] - q2["operating_income_usd"],
                ("Quarter bridge", "B17"): 0,
                ("Backlog commitments", "B10"): a["contracted_backlog_usd"] / a["revenue_usd"],
                ("Backlog commitments", "B14"): 0,
            }
        for (sheet, coordinate), value in expected.items():
            observed = values[sheet][coordinate].value
            assert isinstance(observed, (int, float)), (item["id"], sheet, coordinate, observed)
            assert abs(observed - value) <= max(0.000001, abs(value) * 0.00000001), (item["id"], sheet, coordinate, observed, value)
        checks = [(ws.title, c.coordinate, c.value) for ws in values for row in ws for c in row
                  if isinstance(c.value, str) and c.value in ("RECONCILED", "REVIEW")]
        assert len(checks) >= 5 and all(v == "RECONCILED" for _, _, v in checks)
    else:
        assert values["Requirements"].max_row == 23
        ids = [values["Requirements"].cell(r, 1).value for r in range(8, 24)]
        assert len(set(ids)) == 16
    formulas.close()
    values.close()
    return {"status": "passed", "engine": "LibreOffice 26.8.0.3 native Windows headless import/export",
            "formula_count": formula_count, "formula_errors": 0,
            "expected_values_checked": 13 if item["type"] != "product_spec" else 0,
            "note": "Specification workbooks contain source requirements, not financial formulas." if item["type"] == "product_spec" else "All formulas retain cached results; income, allocation, ratio, and trend expectations checked."}


# Recalculate via the native Windows engine, using a run-isolated profile.
pending_xlsx = [d for d in items if d["format"] == "xlsx"
                and not (checks_dir / f'{d["id"]}.json').exists()]
if pending_xlsx:
    run_lo(["--convert-to", "xlsx:Calc MS Excel 2007 XML", "--outdir", str(recalculated)] +
           [str(ROOT / "documents" / d["filename"]) for d in pending_xlsx])
    for item in pending_xlsx:
        converted = recalculated / item["filename"]
        assert converted.exists(), ("Recalculation missing", item["id"])
        draft = company.load(BUILD / "drafts" / f'{item["id"]}.json')
        check_workbook(item, draft, converted)
        os.replace(converted, ROOT / "documents" / item["filename"])

pending_pdf = [d for d in items if not (renders / d["id"] / "document.pdf").exists()]
for fmt in ("docx", "pptx", "xlsx"):
    group = [d for d in pending_pdf if d["format"] == fmt]
    if not group:
        continue
    temp = renders / f"converted-{fmt}"
    temp.mkdir(exist_ok=True)
    run_lo(["--convert-to", "pdf", "--outdir", str(temp)] +
           [str(ROOT / "documents" / d["filename"]) for d in group])
    for item in group:
        pdf = temp / (Path(item["filename"]).stem + ".pdf")
        assert pdf.exists(), ("Render missing", item["id"])
        folder = renders / item["id"]
        folder.mkdir(exist_ok=True)
        os.replace(pdf, folder / "document.pdf")

for item in items:
    path = ROOT / "documents" / item["filename"]
    draft = company.load(BUILD / "drafts" / f'{item["id"]}.json')
    check_office_text(item, draft, path)
    recalc = check_workbook(item, draft, path) if item["format"] == "xlsx" else None
    folder = renders / item["id"]
    pdf = fitz.open(folder / "document.pdf")
    issues, page_texts = [], []
    thumbs = []
    for n, page in enumerate(pdf):
        text = page.get_text()
        page_texts.append(text)
        for block in page.get_text("dict")["blocks"]:
            for line in block.get("lines", []):
                for span in line["spans"]:
                    x0, y0, x1, y1 = span["bbox"]
                    if x0 < -0.5 or y0 < -0.5 or x1 > page.rect.width + 0.5 or y1 > page.rect.height + 0.5:
                        issues.append({"page": n + 1, "issue": "Text outside page", "text": span["text"]})
        if item["format"] == "pptx":
            for body in draft["slides"][n]["body"]:
                if normalize(body) not in normalize(text):
                    issues.append({"page": n + 1, "issue": "Rendered slide text missing or clipped", "text": body})
        pix = page.get_pixmap(matrix=fitz.Matrix(1.4, 1.4), alpha=False)
        image = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
        image.save(folder / f"page-{n+1:03}.png")
        image.thumbnail((420, 315))
        thumbs.append(image)
    if item["format"] == "docx":
        body_pages = []
        for text in page_texts:
            body_pages.append("\n".join(line for line in text.splitlines()
                                       if line.strip() != "ASTRODIG LLC"
                                       and not line.strip().startswith("AstroDig LLC |")))
        combined = normalize(" ".join(body_pages))
        for section in draft["sections"]:
            if normalize(section["text"]) not in combined:
                issues.append({"issue": "Rendered Word text mismatch", "heading": section["heading"]})
    if item["format"] == "pptx":
        assert len(pdf) == len(draft["slides"]), (item["id"], "Slide count mismatch")
    cols = 3 if item["format"] == "pptx" else 2
    cell_w, cell_h = 440, 345
    sheet = Image.new("RGB", (cols * cell_w, ((len(thumbs) + cols - 1) // cols) * cell_h + 44), "#e4e7e8")
    draw = ImageDraw.Draw(sheet)
    draw.text((12, 10), f'{item["id"]} | {item["type"]} | {item["title"]}', fill="black")
    for n, thumb in enumerate(thumbs):
        x, y = (n % cols) * cell_w + 10, (n // cols) * cell_h + 44
        sheet.paste(thumb, (x, y))
        draw.text((x, y + 316), f"Page {n+1}", fill="black")
    sheet.save(folder / "contact-sheet.png")
    receipt = {"document_id": item["id"], "artifact_sha256": company.file_hash(path),
               "draft_sha256": company.file_hash(BUILD / "drafts" / f'{item["id"]}.json'),
               "pages": len(pdf), "package_integrity": "passed", "extracted_text": "passed",
               "render_text_issues": issues, "recalculation": recalc,
               "measurements": draft["measurements"]}
    (checks_dir / f'{item["id"]}.json').write_text(json.dumps(receipt, indent=2) + "\n")
    pdf.close()
    print(item["id"], "pages", receipt["pages"], "render issues", len(issues),
          "formulas", recalc["formula_count"] if recalc else "-")
