"""Isolated, contract-driven DOCX/XLSX pilot using headless LibreOffice."""

from __future__ import annotations

import argparse
from copy import copy
import hashlib
import importlib.util
import json
import math
import os
import re
import shutil
import subprocess
import tempfile
from pathlib import Path
from typing import Any

from docx import Document
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from pypdf import PdfReader
from pypdf.errors import PdfReadError


FORMULA_ERRORS = {
    "#BLOCKED!",
    "#BUSY!",
    "#CALC!",
    "#CIRCULAR!",
    "#CONNECT!",
    "#DATA!",
    "#DIV/0!",
    "#FIELD!",
    "#GETTING_DATA",
    "#NAME?",
    "#N/A",
    "#NULL!",
    "#NUM!",
    "#REF!",
    "#SPILL!",
    "#UNKNOWN!",
    "#VALUE!",
}
REQUEST = {
    "company_name": "Lumenvale Systems",
    "industry": "analytics software",
    "start_year": 2022,
    "years": 1,
    "documents": 2,
    "seed": 42,
    "content_depth": "light",
    "document_types": ["product_spec"],
    "formats": ["docx", "xlsx"],
}


class BackendError(RuntimeError):
    """An authoring, rendering, contract, or recalculation operation failed."""


def _company_contract():
    helper_path = next(
        (
            candidate
            for candidate in (
                Path(__file__).with_name("company_contract.py"),
                Path("/app/company_contract.py"),
            )
            if candidate.is_file()
        ),
        None,
    )
    if helper_path is None:
        raise BackendError(
            "The repository company-content contract helper is missing. "
            "Run this prototype with its Docker Compose or Dev Container "
            "configuration."
        )
    spec = importlib.util.spec_from_file_location("company_contract", helper_path)
    if spec is None or spec.loader is None:
        raise BackendError(f"Could not load contract helper at {helper_path}.")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _ensure_new_file(path: Path) -> None:
    if path.exists():
        raise BackendError(f"Refusing to overwrite existing output: {path}")
    path.parent.mkdir(parents=True, exist_ok=True)


def _ensure_empty_output_dir(path: Path) -> None:
    if path.is_symlink():
        raise BackendError(f"Output root must not be a symbolic link: {path}")
    if path.exists() and not path.is_dir():
        raise BackendError(f"Output root is not a directory: {path}")
    path.mkdir(parents=True, exist_ok=True)
    if next(path.iterdir(), None) is not None:
        raise BackendError(
            f"Refusing to overwrite a non-empty output directory: {path}. "
            "Choose a new output directory."
        )


def _require_non_root() -> None:
    if hasattr(os, "geteuid") and os.geteuid() == 0:
        raise BackendError("This prototype refuses to run as root.")


def _save_json(contract: Any, path: Path, value: Any) -> None:
    try:
        contract.save(path, value)
    except (OSError, ValueError, contract.Invalid) as exc:
        raise BackendError(f"Could not write contract artifact {path}: {exc}") from exc


def _write_docx(path: Path, draft: dict[str, Any], context: dict[str, Any]) -> None:
    _ensure_new_file(path)
    document = Document()
    document.core_properties.title = draft["title"]
    document.core_properties.author = context["company"]["data"]["name"]
    document.add_heading(draft["title"], level=0)
    document.add_paragraph(draft["publication_date"])

    for section in draft["sections"]:
        document.add_heading(section["heading"], level=1)
        if section["role"] == "requirements":
            for line in section["text"].splitlines():
                document.add_paragraph(line, style="List Bullet")
        elif section["role"] == "acceptance_criteria":
            for line in section["text"].splitlines():
                document.add_paragraph(line, style="List Number")
        else:
            document.add_paragraph(section["text"])
    with path.open("xb") as destination:
        document.save(destination)


def _set_print_layout(sheet) -> None:
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_margins.left = 0.35
    sheet.page_margins.right = 0.35
    sheet.page_margins.top = 0.5
    sheet.page_margins.bottom = 0.5
    sheet.print_options.horizontalCentered = True
    sheet.sheet_view.showGridLines = False
    for row in sheet.iter_rows():
        for cell in row:
            font = copy(cell.font)
            font.name = "Liberation Sans"
            cell.font = font


def _write_xlsx(path: Path, draft: dict[str, Any]) -> dict[str, str]:
    _ensure_new_file(path)
    workbook = Workbook()
    overview = workbook.active
    overview.title = "Specification"
    overview.append([draft["title"]])
    overview.append(["Publication date", draft["publication_date"]])
    overview.append(["Scope", draft["sections"][0]["text"]])
    overview["A1"].font = Font(name="Liberation Sans", size=16, bold=True, color="FFFFFF")
    overview["A1"].fill = PatternFill("solid", fgColor="25445C")
    overview["A1"].alignment = Alignment(
        horizontal="left", vertical="center", indent=1
    )
    overview.merge_cells("A1:B1")
    overview.row_dimensions[1].height = 34
    overview.row_dimensions[2].height = 22
    overview.column_dimensions["A"].width = 26
    overview.column_dimensions["B"].width = 84
    overview["A3"].alignment = Alignment(wrap_text=True, vertical="top")
    overview["B3"].alignment = Alignment(wrap_text=True, vertical="top")
    overview.row_dimensions[3].height = 72
    overview.print_area = "A1:B3"
    overview.freeze_panes = "A2"
    _set_print_layout(overview)

    requirements = workbook.create_sheet("Requirements")
    requirements.append([draft["title"] + " - Requirements"])
    requirements.merge_cells("A1:C1")
    requirements["A1"].font = Font(
        name="Liberation Sans", size=16, bold=True, color="FFFFFF"
    )
    requirements["A1"].fill = PatternFill("solid", fgColor="25445C")
    requirements["A1"].alignment = Alignment(
        horizontal="left", vertical="center", indent=1
    )
    requirements.row_dimensions[1].height = 34
    requirements.append(["Requirement ID", "Category", "Requirement"])
    for cell in requirements[2]:
        cell.font = Font(name="Liberation Sans", bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="25445C")
    for line in draft["sections"][1]["text"].splitlines():
        identifier, category, statement = line.split(" | ", maxsplit=2)
        requirements.append([identifier, category, statement])
    for column, width in {"A": 25, "B": 20, "C": 86}.items():
        requirements.column_dimensions[column].width = width
    for row in requirements.iter_rows(min_row=2):
        row[2].alignment = Alignment(wrap_text=True, vertical="top")
    requirements.freeze_panes = "A3"
    requirements.print_title_rows = "1:2"
    requirements.print_area = f"A1:C{requirements.max_row}"
    _set_print_layout(requirements)

    criteria = workbook.create_sheet("Acceptance Criteria")
    criteria.append([draft["title"] + " - Acceptance Criteria"])
    criteria.merge_cells("A1:B1")
    criteria["A1"].font = Font(
        name="Liberation Sans", size=16, bold=True, color="FFFFFF"
    )
    criteria["A1"].fill = PatternFill("solid", fgColor="25445C")
    criteria["A1"].alignment = Alignment(
        horizontal="left", vertical="center", indent=1
    )
    criteria.row_dimensions[1].height = 34
    criteria.append(["Requirement ID", "Acceptance Criterion"])
    for cell in criteria[2]:
        cell.font = Font(name="Liberation Sans", bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="25445C")
    for line in draft["sections"][2]["text"].splitlines():
        identifier, criterion = line.split(" | ", maxsplit=1)
        criteria.append([identifier, criterion])
    criteria.column_dimensions["A"].width = 25
    criteria.column_dimensions["B"].width = 86
    for row in criteria.iter_rows(min_row=2):
        row[1].alignment = Alignment(wrap_text=True, vertical="top")
    criteria.freeze_panes = "A3"
    criteria.print_title_rows = "1:2"
    criteria.print_area = f"A1:B{criteria.max_row}"
    _set_print_layout(criteria)

    summary = workbook.create_sheet("Validation")
    summary.append(["Validation Summary"])
    summary.merge_cells("A1:B1")
    summary["A1"].font = Font(
        name="Liberation Sans", size=16, bold=True, color="FFFFFF"
    )
    summary["A1"].fill = PatternFill("solid", fgColor="25445C")
    summary["A1"].alignment = Alignment(
        horizontal="left", vertical="center", indent=1
    )
    summary.row_dimensions[1].height = 34
    summary.append(["Validation metric", "Count"])
    for cell in summary[2]:
        cell.font = Font(name="Liberation Sans", bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="25445C")
    summary.append(["Distinct requirements", ""])
    summary.append(["Acceptance criteria", ""])
    summary["B3"] = f"=COUNTA(Requirements!A3:A{requirements.max_row})"
    summary["B4"] = f"=COUNTA('Acceptance Criteria'!B3:B{criteria.max_row})"
    summary.column_dimensions["A"].width = 30
    summary.column_dimensions["B"].width = 18
    summary.print_area = "A1:B4"
    _set_print_layout(summary)

    formulas = {
        "Validation!B3": f"=COUNTA(Requirements!A3:A{requirements.max_row})",
        "Validation!B4": f"=COUNTA('Acceptance Criteria'!B3:B{criteria.max_row})",
    }
    with path.open("xb") as destination:
        workbook.save(destination)
    workbook.close()
    return formulas


def _copy_no_clobber(source: Path, destination: Path) -> None:
    _ensure_new_file(destination)
    with source.open("rb") as source_stream, destination.open("xb") as output_stream:
        shutil.copyfileobj(source_stream, output_stream)
    shutil.copystat(source, destination)


def _convert_with_libreoffice(
    source: Path, output_dir: Path, target_format: str
) -> Path:
    executable = shutil.which("soffice")
    if executable is None:
        raise BackendError("LibreOffice 'soffice' is not installed or is not on PATH.")
    if not source.is_file():
        raise BackendError(f"Input file does not exist: {source}")
    output_dir.mkdir(parents=True, exist_ok=True)

    with tempfile.TemporaryDirectory(prefix="lo-user-profile-") as profile_dir:
        command = [
            executable,
            "--headless",
            "--nologo",
            "--nodefault",
            "--nolockcheck",
            "--norestore",
            f"-env:UserInstallation={Path(profile_dir).as_uri()}",
            "--convert-to",
            target_format,
            "--outdir",
            str(output_dir),
            str(source),
        ]
        try:
            result = subprocess.run(
                command,
                check=False,
                capture_output=True,
                text=True,
                timeout=120,
            )
        except subprocess.TimeoutExpired as exc:
            raise BackendError(
                f"LibreOffice timed out converting {source.name} to {target_format}."
            ) from exc

    converted = output_dir / f"{source.stem}.{target_format}"
    if result.returncode != 0 or not converted.is_file():
        details = "\n".join(
            part.strip() for part in (result.stdout, result.stderr) if part.strip()
        )
        raise BackendError(
            f"LibreOffice failed to convert {source.name} to {target_format}."
            + (f"\n{details}" if details else "")
        )
    return converted


def validate_pdf(path: Path, expected_text: tuple[str, ...] = ()) -> dict[str, Any]:
    if not path.is_file():
        raise BackendError(f"Expected PDF was not produced: {path}")
    pdf_bytes = path.read_bytes()
    if not pdf_bytes.startswith(b"%PDF-") or b"%%EOF" not in pdf_bytes[-1024:]:
        raise BackendError(f"PDF header or end marker is missing: {path}")
    try:
        reader = PdfReader(path, strict=True)
        if reader.is_encrypted:
            raise BackendError(f"Rendered PDF is unexpectedly encrypted: {path}")
        pages = list(reader.pages)
        page_text = [page.extract_text() or "" for page in pages]
    except (OSError, ValueError, PdfReadError) as exc:
        raise BackendError(f"Could not parse rendered PDF {path}: {exc}") from exc
    if not pages:
        raise BackendError(f"Rendered PDF has no pages: {path}")
    if any(not text.strip() for text in page_text):
        raise BackendError(f"Rendered PDF contains a page with no extractable text: {path}")
    complete_text = re.sub(r"\s+", " ", " ".join(page_text)).casefold()
    for expected in expected_text:
        if re.sub(r"\s+", " ", expected).casefold() not in complete_text:
            raise BackendError(
                f"Expected text is missing from rendered PDF {path}: {expected!r}"
            )
    return {"pages": len(pages), "text_characters": len(complete_text)}


def _rasterize_pdf(path: Path, output_dir: Path, stem: str, pages: int) -> list[Path]:
    executable = shutil.which("pdftoppm")
    if executable is None:
        raise BackendError("PDF visual inspection requires the pinned pdftoppm utility.")
    output_dir.mkdir(parents=True, exist_ok=True)
    images = []
    for page_number in range(1, pages + 1):
        image_stem = output_dir / f"{stem}-page-{page_number}"
        image = image_stem.with_suffix(".png")
        _ensure_new_file(image)
        try:
            result = subprocess.run(
                [
                    executable,
                    "-f",
                    str(page_number),
                    "-l",
                    str(page_number),
                    "-singlefile",
                    "-png",
                    "-r",
                    "110",
                    str(path),
                    str(image_stem),
                ],
                check=False,
                capture_output=True,
                text=True,
                timeout=60,
            )
        except subprocess.TimeoutExpired as exc:
            raise BackendError(
                f"PDF page rasterization timed out: {path}, page {page_number}"
            ) from exc
        if result.returncode != 0 or not image.is_file():
            details = result.stderr.strip() or result.stdout.strip()
            raise BackendError(
                f"Could not rasterize {path}, page {page_number}."
                + (f"\n{details}" if details else "")
            )
        if not image.read_bytes().startswith(b"\x89PNG\r\n\x1a\n"):
            raise BackendError(f"Rasterized output is not a PNG image: {image}")
        images.append(image)
    return images


def render_docx_to_pdf(
    source: Path, destination: Path, expected_text: tuple[str, ...] = ()
) -> dict[str, Any]:
    if source.suffix.lower() != ".docx":
        raise BackendError(f"Expected a DOCX input, got: {source.name}")
    with tempfile.TemporaryDirectory(prefix="lo-pdf-") as temporary_dir:
        converted = _convert_with_libreoffice(source, Path(temporary_dir), "pdf")
        _copy_no_clobber(converted, destination)
    return validate_pdf(destination, expected_text)


def render_xlsx_to_pdf(
    source: Path, destination: Path, expected_text: tuple[str, ...] = ()
) -> dict[str, Any]:
    if source.suffix.lower() != ".xlsx":
        raise BackendError(f"Expected an XLSX input, got: {source.name}")
    with tempfile.TemporaryDirectory(prefix="lo-pdf-") as temporary_dir:
        converted = _convert_with_libreoffice(source, Path(temporary_dir), "pdf")
        _copy_no_clobber(converted, destination)
    return validate_pdf(destination, expected_text)


def _read_formulas(path: Path) -> dict[str, str]:
    workbook = load_workbook(path, data_only=False)
    try:
        formulas = {}
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows():
                for cell in row:
                    if cell.data_type == "f":
                        formulas[f"{sheet.title}!{cell.coordinate}"] = str(cell.value)
        return formulas
    finally:
        workbook.close()


def recalculate_xlsx(source: Path, destination: Path) -> Path:
    if source.suffix.lower() != ".xlsx":
        raise BackendError(f"Expected an XLSX input, got: {source.name}")
    _ensure_new_file(destination)
    expected_formulas = _read_formulas(source)
    with tempfile.TemporaryDirectory(prefix="lo-xlsx-") as temporary_dir:
        converted = _convert_with_libreoffice(source, Path(temporary_dir), "xlsx")
        _copy_no_clobber(converted, destination)
    actual_formulas = _read_formulas(destination)
    if actual_formulas != expected_formulas:
        raise BackendError(
            "LibreOffice changed the workbook's formula cells or expressions: "
            f"expected {expected_formulas}, found {actual_formulas}"
        )
    return destination


def inspect_recalculated_workbook(
    path: Path,
    expected_values: dict[str, float] | None = None,
    expected_formulas: dict[str, str] | None = None,
) -> dict[str, dict[str, Any]]:
    formulas_book = load_workbook(path, data_only=False)
    cached_book = load_workbook(path, data_only=True)
    formulas: dict[str, str] = {}
    cached_values: dict[str, Any] = {}
    errors: list[str] = []
    missing_values: list[str] = []
    try:
        for formula_sheet in formulas_book.worksheets:
            cached_sheet = cached_book[formula_sheet.title]
            for row in formula_sheet.iter_rows():
                for formula_cell in row:
                    if formula_cell.data_type != "f":
                        continue
                    reference = f"{formula_sheet.title}!{formula_cell.coordinate}"
                    formulas[reference] = str(formula_cell.value)
                    cached_cell = cached_sheet[formula_cell.coordinate]
                    if cached_cell.value is None:
                        missing_values.append(reference)
                    else:
                        cached_values[reference] = cached_cell.value

        for cached_sheet in cached_book.worksheets:
            for row in cached_sheet.iter_rows():
                for cell in row:
                    if cell.data_type == "e" or (
                        isinstance(cell.value, str)
                        and cell.value.upper() in FORMULA_ERRORS
                    ):
                        errors.append(
                            f"{cached_sheet.title}!{cell.coordinate}={cell.value}"
                        )
    finally:
        formulas_book.close()
        cached_book.close()

    if errors:
        raise BackendError("Formula errors found: " + ", ".join(errors))
    if not formulas:
        raise BackendError("No formulas were found in the recalculated workbook.")
    if missing_values:
        raise BackendError(
            "Formula results are missing from cached values: "
            + ", ".join(missing_values)
        )
    if expected_formulas is not None and formulas != expected_formulas:
        raise BackendError(
            f"Formula expressions differ: expected {expected_formulas}, got {formulas}"
        )

    if expected_values:
        for reference, expected in expected_values.items():
            if reference not in cached_values:
                raise BackendError(f"Expected formula result is missing: {reference}")
            actual = cached_values[reference]
            if (
                isinstance(actual, bool)
                or not isinstance(actual, (int, float))
                or not math.isfinite(actual)
                or not math.isclose(actual, expected, rel_tol=1e-9, abs_tol=1e-9)
            ):
                raise BackendError(
                    f"Unexpected cached result for {reference}: "
                    f"expected {expected}, got {actual!r}"
                )

    return {"formulas": formulas, "cached_values": cached_values}


def _add_requirement_facts(contract: Any, canon: dict[str, Any]) -> None:
    for product_id, requirement_id, category, statement, criteria in (
        (
            "product-beacon",
            "BEACON-1.0-FR-01",
            "Functional",
            "Store each observation against a unique mission-profile identifier.",
            [
                "Saving two mission profiles preserves distinct identifiers.",
                "A reopened profile retains its original identifier.",
            ],
        ),
        (
            "product-beacon",
            "BEACON-1.0-NFR-01",
            "Reliability",
            "Retain each accepted observation record when a session is reopened.",
            [
                "A reopened session presents every previously accepted observation.",
                "The retained record count matches the saved session count.",
            ],
        ),
        (
            "product-prism",
            "PRISM-1.0-FR-01",
            "Functional",
            "Associate each analysis result with its source observation identifier.",
            [
                "An analysis result displays the identifier of its source observation.",
                "Results without a source identifier are rejected.",
            ],
        ),
        (
            "product-prism",
            "PRISM-1.0-NFR-01",
            "Reliability",
            "Preserve completed analysis results across a normal application restart.",
            [
                "A completed result remains available after a normal restart.",
                "The result values before and after restart are identical.",
            ],
        ),
    ):
        canon["records"].append(
            {
                "id": "requirement-" + requirement_id.lower(),
                "kind": "requirement",
                "known_on": "2022-01-01",
                "effective_on": "2022-01-01",
                "data": {
                    "product_id": product_id,
                    "version": "1.0",
                    "requirement_id": requirement_id,
                    "category": category,
                    "statement": statement,
                    "acceptance_criteria": criteria,
                },
                "refs": [product_id],
            }
        )
    contract.validate_canon(canon)


def _draft_from_context(contract: Any, context: dict[str, Any]) -> dict[str, Any]:
    document = context["document"]
    records = {record["id"]: record for record in context["records"]}
    anchor = records[document["anchor_ids"][0]]
    product = records[anchor["data"]["product_id"]]
    requirements = [
        record
        for record in context["records"]
        if record["kind"] == "requirement"
        and record["data"]["product_id"] == anchor["data"]["product_id"]
        and record["data"]["version"] == anchor["data"]["version"]
    ]
    if not requirements:
        raise BackendError(f"No date-visible requirement facts for {document['id']}.")

    claims = []
    claim_ids: dict[tuple[str, tuple[Any, ...]], str] = {}

    def claim(record_id: str, path: list[Any]) -> tuple[str, Any]:
        key = (record_id, tuple(path))
        if key not in claim_ids:
            identifier = f"claim-{len(claims) + 1:03d}"
            claim_ids[key] = identifier
            record = records[record_id]
            claims.append(
                {
                    "id": identifier,
                    "record_id": record_id,
                    "path": path,
                    "value": contract.resolve(record, path),
                }
            )
        return claim_ids[key], contract.resolve(records[record_id], path)

    company_claim, company_name = claim(context["company"]["id"], ["data", "name"])
    product_claim, product_name = claim(product["id"], ["data", "name"])
    version_claim, version = claim(anchor["id"], ["data", "version"])
    status_claim, status = claim(anchor["id"], ["data", "status"])
    target_claim, target_date = claim(anchor["id"], ["data", "target_date"])

    section_claims = [
        company_claim,
        product_claim,
        version_claim,
        status_claim,
        target_claim,
    ]
    sections = [
        {
            "role": "scope",
            "heading": "Scope and release status",
            "text": (
                f"{company_name} is preparing the {product_name} {version} "
                f"specification. Its canonical release milestone is {status}, "
                f"with a target date of {target_date}; this record does not state "
                "that the release has launched."
            ),
            "claim_ids": section_claims,
        }
    ]

    requirement_lines = []
    requirement_claims = []
    criteria_lines = []
    criteria_claims = []
    for record in requirements:
        data = record["data"]
        requirement_id_claim, requirement_id = claim(
            record["id"], ["data", "requirement_id"]
        )
        category_claim, category = claim(record["id"], ["data", "category"])
        statement_claim, statement = claim(record["id"], ["data", "statement"])
        requirement_claims.extend(
            [requirement_id_claim, category_claim, statement_claim]
        )
        requirement_lines.append(f"{requirement_id} | {category} | {statement}")
        for criterion_index, criterion in enumerate(data["acceptance_criteria"]):
            criterion_claim, criterion_value = claim(
                record["id"], ["data", "acceptance_criteria", criterion_index]
            )
            criteria_claims.extend([requirement_id_claim, criterion_claim])
            criteria_lines.append(f"{requirement_id} | {criterion_value}")

    sections.extend(
        [
            {
                "role": "requirements",
                "heading": "Requirements",
                "text": "\n".join(requirement_lines),
                "claim_ids": requirement_claims,
            },
            {
                "role": "acceptance_criteria",
                "heading": "Acceptance criteria",
                "text": "\n".join(criteria_lines),
                "claim_ids": criteria_claims,
            },
        ]
    )
    return {
        "document_id": document["id"],
        "canon_sha256": context["canon_sha256"],
        "plan_sha256": context["plan_sha256"],
        "title": document["title"],
        "publication_date": document["publication_date"],
        "content_depth": document["content_depth"],
        "claims": claims,
        "sections": sections,
    }


def run_demo(output_dir: Path) -> dict[str, Any]:
    _require_non_root()
    contract = _company_contract()
    _ensure_empty_output_dir(output_dir)
    build_dir = output_dir / "build"
    documents_dir = output_dir / "documents"
    contexts_dir = build_dir / "contexts"
    drafts_dir = build_dir / "drafts"
    renders_dir = build_dir / "renders"
    for directory in (build_dir, documents_dir, contexts_dir, drafts_dir, renders_dir):
        directory.mkdir()

    canon = contract.seed_canon(REQUEST)
    _add_requirement_facts(contract, canon)
    plan = contract.create_plan(canon)
    by_format = {document["format"]: document for document in plan["documents"]}
    if set(by_format) != {"docx", "xlsx"} or len(plan["documents"]) != 2:
        raise BackendError(f"Expected one DOCX and one XLSX plan item, got {plan}.")

    _save_json(contract, output_dir / "canon.json", canon)
    _save_json(contract, build_dir / "request.json", canon["request"])
    _save_json(contract, build_dir / "plan.json", plan)
    contract.validate_output_layout(canon, output_dir)

    pilot_documents = []
    recalculated_results: dict[str, Any] | None = None
    expected_formulas: dict[str, str] | None = None
    for document in plan["documents"]:
        context = contract.get_context(canon, plan, document["id"])
        draft = _draft_from_context(contract, context)
        contract.check_draft(canon, plan, draft)
        context_path = contexts_dir / f"{document['id']}.json"
        draft_path = drafts_dir / f"{document['id']}.json"
        _save_json(contract, context_path, context)
        _save_json(contract, draft_path, draft)
        artifact_path = documents_dir / document["filename"]
        if document["format"] == "docx":
            _write_docx(artifact_path, draft, context)
            contract.audit_package(artifact_path, "docx")
            pdf_path = renders_dir / f"{artifact_path.stem}.pdf"
            pdf_evidence = render_docx_to_pdf(
                artifact_path,
                pdf_path,
                (
                    document["title"],
                    "Scope and release status",
                    draft["sections"][1]["text"].split(" | ")[0],
                ),
            )
        else:
            expected_formulas = _write_xlsx(artifact_path, draft)
            contract.audit_package(artifact_path, "xlsx")
            recalculated_path = renders_dir / f"recalculated-{artifact_path.name}"
            recalculate_xlsx(artifact_path, recalculated_path)
            contract.audit_package(recalculated_path, "xlsx")
            expected_values = {
                "Validation!B3": float(len(draft["sections"][1]["text"].splitlines())),
                "Validation!B4": float(len(draft["sections"][2]["text"].splitlines())),
            }
            recalculated_results = inspect_recalculated_workbook(
                recalculated_path,
                expected_values=expected_values,
                expected_formulas=expected_formulas,
            )
            _copy_no_clobber(recalculated_path, renders_dir / artifact_path.name)
            pdf_path = renders_dir / f"{artifact_path.stem}.pdf"
            pdf_evidence = render_xlsx_to_pdf(
                recalculated_path,
                pdf_path,
                (
                    document["title"],
                    "Distinct requirements",
                    "Acceptance criteria",
                    draft["sections"][1]["text"].split(" | ")[0],
                ),
            )
        images = _rasterize_pdf(
            pdf_path,
            renders_dir / "visual",
            document["id"],
            pdf_evidence["pages"],
        )
        pilot_documents.append(
            {
                "document_id": document["id"],
                "format": document["format"],
                "title": document["title"],
                "publication_date": document["publication_date"],
                "context_sha256": _sha256(context_path),
                "draft_sha256": _sha256(draft_path),
                "artifact": artifact_path.relative_to(output_dir).as_posix(),
                "artifact_sha256": _sha256(artifact_path),
                "pdf": pdf_path.relative_to(output_dir).as_posix(),
                "pdf_validation": pdf_evidence,
                "visual_pages": [
                    image.relative_to(output_dir).as_posix() for image in images
                ],
            }
        )

    contract.check_drafts(canon, plan, drafts_dir)
    if recalculated_results is None or expected_formulas is None:
        raise BackendError("The planned XLSX recalculation step did not run.")
    results = {
        "schema_version": 1,
        "scope": "prototype smoke evidence only; not production QA or Office fidelity certification",
        "contract_validation": {
            "planned_documents": len(plan["documents"]),
            "formats": sorted({document["format"] for document in plan["documents"]}),
            "checked_drafts": len(plan["documents"]),
            "date_filtered_contexts": True,
            "office_package_integrity": True,
        },
        "documents": pilot_documents,
        "recalculation": recalculated_results,
        "office_compatibility_qa": "not performed",
        "depth_review": "not performed; pilot requirements are not a full genre-depth deliverable",
    }
    _save_json(contract, build_dir / "pilot-results.json", results)
    print(json.dumps(results, indent=2))
    return results


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Generate and validate a contract-driven synthetic DOCX/XLSX pilot."
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=Path("/output"),
        help="empty output directory or new child under the mounted output directory",
    )
    args = parser.parse_args()
    run_demo(args.output)


if __name__ == "__main__":
    main()
