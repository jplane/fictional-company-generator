"""Export prepared company canon/plan/context/draft artifacts in the backend container."""

from __future__ import annotations

import argparse
import json
import math
import os
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches as DocxInches
from docx.shared import Pt as DocxPt
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Inches, Pt

import office_backend as office


ROOT_ENTRIES = {"canon.json", "documents", "build"}
PRIVATE_DRAFT_KEYS = {"artifact", "claims", "measurements"}
COMMON_DRAFT_KEYS = {
    "document_id",
    "canon_sha256",
    "plan_sha256",
    "title",
    "publication_date",
    "content_depth",
    "claims",
    "sections",
    "artifact",
    "measurements",
    "slides",
    "model",
    "requirement_rows",
}
SECTION_KEYS = {"role", "heading", "text", "claim_ids", "tables"}
SLIDE_KEYS = {
    "heading",
    "title",
    "subtitle",
    "body",
    "layout",
    "badge",
    "chart",
    "tables",
    "notes",
}
TABLE_KEYS = {"title", "headers", "rows"}
PPT_LAYOUTS = {
    "capabilities",
    "cards",
    "chart",
    "closing",
    "content",
    "cover",
    "grid",
    "statement",
    "stats",
    "workflow",
}
CHART_TYPES = {
    "bar": XL_CHART_TYPE.BAR_CLUSTERED,
    "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
    "doughnut": XL_CHART_TYPE.DOUGHNUT,
    "line": XL_CHART_TYPE.LINE_MARKERS,
    "pie": XL_CHART_TYPE.PIE,
}
FORMULA_ERRORS = office.FORMULA_ERRORS
NAVY = (37, 68, 92)
BLUE = (34, 115, 145)
TEAL = (35, 139, 132)
PALE = (233, 242, 245)
INK = (44, 55, 64)
MUTED = (105, 118, 126)
WHITE = (255, 255, 255)


def _fail(message: str) -> None:
    raise office.BackendError(message)


def _json_file(path: Path) -> Any:
    if path.is_symlink() or not path.is_file():
        _fail(f"Expected a regular JSON file, not a link: {path}")
    try:
        with path.open(encoding="utf-8") as stream:
            return json.load(stream)
    except (OSError, UnicodeError, json.JSONDecodeError) as exc:
        _fail(f"Could not read JSON input {path}: {exc}")


def _write_json_new(path: Path, value: Any) -> None:
    if path.exists() or path.is_symlink():
        _fail(f"Refusing to overwrite existing output: {path}")
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("x", encoding="utf-8", newline="\n") as stream:
        json.dump(value, stream, indent=2, ensure_ascii=False, allow_nan=False)
        stream.write("\n")


def _sha256(path: Path) -> str:
    return office._sha256(path)


def _reject_link(path: Path, label: str) -> None:
    if path.is_symlink():
        _fail(f"{label} must not be a symbolic link: {path}")


def _ensure_directory(path: Path, label: str) -> None:
    _reject_link(path, label)
    if path.exists() and not path.is_dir():
        _fail(f"{label} is not a directory: {path}")
    path.mkdir(parents=True, exist_ok=True)


def _same_input(source: Path, destination: Path, contract: Any) -> bool:
    if not destination.is_file() or destination.is_symlink():
        return False
    if source.suffix.casefold() == ".json" and destination.suffix.casefold() == ".json":
        try:
            return contract.fingerprint(_json_file(source)) == contract.fingerprint(
                _json_file(destination)
            )
        except (office.BackendError, contract.Invalid, TypeError, ValueError):
            return False
    return _sha256(source) == _sha256(destination)


def _stage_file(source: Path, destination: Path, contract: Any) -> None:
    _reject_link(destination, "Staged input")
    if destination.exists():
        if _same_input(source, destination, contract):
            return
        _fail(f"Output corpus has a different file at the input path: {destination}")
    destination.parent.mkdir(parents=True, exist_ok=True)
    with source.open("rb") as input_stream, destination.open("xb") as output_stream:
        shutil.copyfileobj(input_stream, output_stream)


def _json_paths(directory: Path, label: str, planned_ids: set[str]) -> dict[str, Path]:
    _reject_link(directory, label)
    if not directory.is_dir():
        _fail(f"{label} directory does not exist: {directory}")
    paths: dict[str, Path] = {}
    for path in directory.glob("*.json"):
        _reject_link(path, f"{label} input")
        if not path.is_file():
            _fail(f"Invalid {label} input: {path}")
        identifier = path.stem
        if identifier not in planned_ids:
            _fail(f"{label} contains an ID not present in the plan: {path.name}")
        if identifier in paths:
            _fail(f"Duplicate {label} ID: {identifier}")
        paths[identifier] = path
    return paths


def _is_scalar(value: Any) -> bool:
    return value is None or isinstance(value, (str, int, float, bool))


def _normalize(value: str) -> str:
    return re.sub(r"\s+", " ", value).strip().casefold()


def _validate_table(table: Any, label: str, fmt: str) -> None:
    if not isinstance(table, dict):
        _fail(f"{label} must be a table object")
    unknown = set(table) - TABLE_KEYS
    if unknown:
        _fail(f"{label} has unsupported table fields: {sorted(unknown)}")
    headers = table.get("headers")
    rows = table.get("rows")
    if (
        not isinstance(headers, list)
        or not headers
        or not all(isinstance(value, str) and value.strip() for value in headers)
    ):
        _fail(f"{label} requires non-empty text headers")
    if not isinstance(rows, list) or not rows:
        _fail(f"{label} requires at least one data row")
    for row_number, row in enumerate(rows, start=1):
        if not isinstance(row, list) or len(row) != len(headers):
            _fail(f"{label} row {row_number} must match its header count")
        for cell in row:
            if _is_scalar(cell):
                continue
            if (
                fmt == "xlsx"
                and isinstance(cell, dict)
                and set(cell) <= {"formula", "expected_value"}
                and isinstance(cell.get("formula"), str)
                and cell["formula"].startswith("=")
                and "[" not in cell["formula"]
                and (
                    "expected_value" not in cell
                    or (
                        isinstance(cell["expected_value"], (int, float))
                        and not isinstance(cell["expected_value"], bool)
                        and math.isfinite(cell["expected_value"])
                    )
                )
            ):
                continue
            _fail(
                f"{label} contains an unsupported cell value; use scalar cells or "
                "an XLSX formula object with an optional numeric expected_value"
            )


def _body_lines(value: Any, label: str) -> list[str]:
    if isinstance(value, str):
        lines = [line for line in value.splitlines() if line.strip()]
    elif isinstance(value, list) and all(isinstance(item, str) for item in value):
        lines = [item for item in value if item.strip()]
    else:
        _fail(f"{label}.body must be text or a list of text")
    if not lines:
        _fail(f"{label}.body must not be empty")
    return lines


def _validate_draft_structure(draft: dict[str, Any], document: dict[str, Any]) -> None:
    fmt = document["format"]
    unknown = set(draft) - COMMON_DRAFT_KEYS
    if unknown:
        _fail(f"{document['id']} draft has unsupported top-level fields: {sorted(unknown)}")
    sections = draft.get("sections")
    if not isinstance(sections, list) or not sections:
        _fail(f"{document['id']} draft requires sections")
    for index, section in enumerate(sections, start=1):
        if not isinstance(section, dict):
            _fail(f"{document['id']} section {index} must be an object")
        extra = set(section) - SECTION_KEYS
        if extra:
            _fail(
                f"{document['id']} section {index} has unsupported structures: "
                f"{sorted(extra)}"
            )
        tables = section.get("tables", [])
        if not isinstance(tables, list):
            _fail(f"{document['id']} section {index}.tables must be a list")
        for table_index, table in enumerate(tables, start=1):
            _validate_table(table, f"{document['id']} section {index} table {table_index}", fmt)

    slides = draft.get("slides")
    if fmt != "pptx" and slides:
        _fail(f"{document['id']} uses a slides structure with non-PPTX format {fmt}")
    if fmt == "pptx":
        if slides is not None and (not isinstance(slides, list) or not slides):
            _fail(f"{document['id']} PPTX slides must be a non-empty list")
        if slides is not None:
            if len(slides) != len(sections):
                _fail(
                    f"{document['id']} has {len(slides)} slides but {len(sections)} "
                    "grounded sections; provide one content-matched slide per section"
                )
            for index, (slide, section) in enumerate(zip(slides, sections), start=1):
                label = f"{document['id']} slide {index}"
                if not isinstance(slide, dict):
                    _fail(f"{label} must be an object")
                extra = set(slide) - SLIDE_KEYS
                if extra:
                    _fail(f"{label} has unsupported structures: {sorted(extra)}")
                heading = slide.get("heading", slide.get("title"))
                if not isinstance(heading, str) or heading != section["heading"]:
                    _fail(f"{label} heading must match its grounded section heading")
                lines = _body_lines(slide.get("body"), label)
                if _normalize("\n".join(lines)) != _normalize(section["text"]):
                    _fail(
                        f"{label} body must contain the full text of its grounded section; "
                        "the exporter will not silently drop or move narrative content"
                    )
                layout = slide.get("layout", "content")
                if layout not in PPT_LAYOUTS:
                    _fail(f"{label} uses unsupported layout: {layout!r}")
                tables = slide.get("tables", [])
                if not isinstance(tables, list):
                    _fail(f"{label}.tables must be a list")
                for table_index, table in enumerate(tables, start=1):
                    _validate_table(
                        table, f"{label} table {table_index}", "pptx"
                    )
                chart = slide.get("chart")
                if chart is not None:
                    if not isinstance(chart, dict) or set(chart) - {
                        "labels",
                        "values",
                        "title",
                        "kind",
                    }:
                        _fail(f"{label} has an unsupported chart structure")
                    labels, values = chart.get("labels"), chart.get("values")
                    if (
                        not isinstance(labels, list)
                        or not labels
                        or not all(isinstance(item, str) and item.strip() for item in labels)
                        or not isinstance(values, list)
                        or len(values) != len(labels)
                        or not all(
                            isinstance(item, (int, float))
                            and not isinstance(item, bool)
                            and math.isfinite(item)
                            for item in values
                        )
                    ):
                        _fail(f"{label} chart labels/values must be matching text/numeric lists")
                    if chart.get("kind", "column") not in CHART_TYPES:
                        _fail(f"{label} uses unsupported chart kind: {chart.get('kind')!r}")
                if chart is not None and (tables or section.get("tables")):
                    _fail(
                        f"{label} cannot combine a chart and table on one slide; "
                        "split them into separate planned sections"
                    )
                notes = slide.get("notes")
                if notes is not None and (
                    not isinstance(notes, str)
                    or _normalize(notes) not in _normalize(section["text"])
                ):
                    _fail(
                        f"{label} notes must be present in its cited section text"
                    )

    model = draft.get("model")
    if model is not None and fmt != "xlsx":
        _fail(f"{document['id']} has structured model data unsupported by {fmt}")
    if model is not None and not isinstance(model, (dict, list)):
        _fail(f"{document['id']} model must be an object or list")
    requirement_rows = draft.get("requirement_rows")
    if requirement_rows is not None:
        if fmt not in {"docx", "xlsx"}:
            _fail(f"{document['id']} requirement_rows are unsupported by {fmt}")
        if not isinstance(requirement_rows, list) or not all(
            isinstance(row, dict) and row for row in requirement_rows
        ):
            _fail(f"{document['id']} requirement_rows must be a list of objects")
        if any(not _is_scalar(value) and not isinstance(value, (list, dict))
               for row in requirement_rows for value in row.values()):
            _fail(f"{document['id']} requirement_rows contain an unsupported value")


def _simple_value(value: Any, label: str) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "Yes" if value else "No"
    if isinstance(value, (int, float)):
        if isinstance(value, float) and not math.isfinite(value):
            _fail(f"{label} must not be NaN or infinite")
        return str(value)
    if isinstance(value, str):
        return value
    _fail(f"{label} must be a scalar value")


def _docx_cell_text(value: Any, label: str) -> str:
    return _simple_value(value, label)


def _add_docx_table(document: Document, table_data: dict[str, Any]) -> None:
    if table_data.get("title"):
        document.add_heading(table_data["title"], level=2)
    table = document.add_table(rows=1, cols=len(table_data["headers"]))
    table.style = "Light Shading Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    header_properties = table.rows[0]._tr.get_or_add_trPr()
    repeating_header = OxmlElement("w:tblHeader")
    repeating_header.set(qn("w:val"), "true")
    header_properties.append(repeating_header)
    for cell, value in zip(table.rows[0].cells, table_data["headers"]):
        cell.text = value
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for run in cell.paragraphs[0].runs:
            run.bold = True
    for row in table_data["rows"]:
        output_row = table.add_row()
        row_properties = output_row._tr.get_or_add_trPr()
        row_properties.append(OxmlElement("w:cantSplit"))
        cells = output_row.cells
        for cell, value in zip(cells, row):
            cell.text = _docx_cell_text(value, "DOCX table cell")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    document.add_paragraph()


def _set_section_orientation(section: Any, landscape: bool) -> None:
    width, height = section.page_width, section.page_height
    if landscape and width < height or not landscape and width > height:
        section.page_width, section.page_height = height, width
    section.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT


def _add_docx_text(document: Document, text: str) -> None:
    for paragraph in re.split(r"\n[ \t]*\n", text.strip()):
        if paragraph.strip():
            document.add_paragraph(paragraph.strip())


def _add_requirement_rows_docx(document: Document, rows: list[dict[str, Any]]) -> None:
    if not rows:
        return
    keys = list(dict.fromkeys(key for row in rows for key in row))
    if any(not isinstance(key, str) or not key.strip() for key in keys):
        _fail("requirement_rows keys must be non-empty text")
    document.add_heading("Requirement traceability", level=1)
    table = document.add_table(rows=1, cols=len(keys))
    table.style = "Light Shading Accent 1"
    for cell, key in zip(table.rows[0].cells, keys):
        cell.text = key.replace("_", " ").title()
        for run in cell.paragraphs[0].runs:
            run.bold = True
    for row in rows:
        cells = table.add_row().cells
        for cell, key in zip(cells, keys):
            value = row.get(key)
            cell.text = (
                json.dumps(value, ensure_ascii=False, sort_keys=True)
                if isinstance(value, (dict, list))
                else _docx_cell_text(value, f"requirement_rows.{key}")
            )


def _write_docx(path: Path, draft: dict[str, Any], context: dict[str, Any]) -> None:
    document = Document()
    properties = document.core_properties
    properties.title = draft["title"]
    properties.author = context["company"]["data"]["name"]
    properties.subject = draft["publication_date"]
    first = document.add_heading(draft["title"], level=0)
    first.alignment = WD_ALIGN_PARAGRAPH.LEFT
    date_paragraph = document.add_paragraph(draft["publication_date"])
    date_paragraph.paragraph_format.space_after = DocxPt(16)

    for section in draft["sections"]:
        document.add_heading(section["heading"], level=1)
        _add_docx_text(document, section["text"])
        for table in section.get("tables", []):
            wide_table = len(table["headers"]) > 5
            if wide_table:
                landscape_section = document.add_section(WD_SECTION.NEW_PAGE)
                _set_section_orientation(landscape_section, True)
            _add_docx_table(document, table)
            if wide_table:
                portrait_section = document.add_section(WD_SECTION.NEW_PAGE)
                _set_section_orientation(portrait_section, False)
    if draft.get("requirement_rows"):
        _add_requirement_rows_docx(document, draft["requirement_rows"])
    for section in document.sections:
        section.top_margin = DocxInches(0.7)
        section.bottom_margin = DocxInches(0.65)
        section.left_margin = DocxInches(0.75)
        section.right_margin = DocxInches(0.75)
    footer = document.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run(context["company"]["data"]["name"] + " | " + draft["publication_date"])
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(path)
    reopened = Document(path)
    actual = "\n".join(
        [paragraph.text for paragraph in reopened.paragraphs]
        + [
            cell.text
            for table in reopened.tables
            for row in table.rows
            for cell in row.cells
        ]
    )
    for expected in [draft["title"], *[s["heading"] for s in draft["sections"]]]:
        if _normalize(expected) not in _normalize(actual):
            _fail(f"DOCX structure check failed to extract: {expected!r}")
    for section in draft["sections"]:
        text_parts = [part for part in re.split(r"\n[ \t]*\n", section["text"].strip()) if part]
        for text_part in text_parts:
            if _normalize(text_part) not in _normalize(actual):
                _fail(
                    f"DOCX structure check failed to extract section text "
                    f"{section['heading']!r}"
                )
    reopened.core_properties  # Force package/property parsing before delivery.


def _sheet_name(value: str, used: set[str]) -> str:
    cleaned = re.sub(r"[\[\]:*?/\\]", " ", value).strip().strip("'") or "Content"
    cleaned = cleaned[:31].rstrip()
    candidate = cleaned
    suffix = 2
    while candidate.casefold() in used:
        tail = f" {suffix}"
        candidate = cleaned[: 31 - len(tail)].rstrip() + tail
        suffix += 1
    used.add(candidate.casefold())
    return candidate


def _style_xlsx_sheet(sheet) -> None:
    sheet.sheet_view.showGridLines = False
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_margins.left = 0.35
    sheet.page_margins.right = 0.35
    sheet.page_margins.top = 0.5
    sheet.page_margins.bottom = 0.5
    sheet.print_options.horizontalCentered = True
    for row in sheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    sheet.freeze_panes = "A2"


def _xlsx_title(sheet, title: str, end_col: int) -> None:
    last = get_column_letter(max(end_col, 1))
    sheet.merge_cells(f"A1:{last}1")
    cell = sheet["A1"]
    cell.value = title
    cell.font = Font(name="Liberation Sans", size=15, bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="25445C")
    cell.alignment = Alignment(vertical="center", indent=1)
    sheet.row_dimensions[1].height = 32


def _xlsx_formula_or_value(
    cell: Any,
    value: Any,
    label: str,
    expected_values: dict[str, float],
) -> None:
    if isinstance(value, dict):
        formula = value["formula"]
        cell.value = formula
        if "expected_value" in value:
            expected_values[f"{cell.parent.title}!{cell.coordinate}"] = float(
                value["expected_value"]
            )
    elif isinstance(value, str) and value.startswith("="):
        cell.value = value
        cell.data_type = "s"
    else:
        cell.value = value
    if isinstance(cell.value, str) and len(cell.value) > 32767:
        _fail(f"{label} exceeds Excel's cell text limit")
    cell.alignment = Alignment(vertical="top", wrap_text=True)


def _xlsx_table(
    sheet: Any,
    table_data: dict[str, Any],
    row_start: int,
    expected_values: dict[str, float],
    label: str,
) -> int:
    headers = table_data["headers"]
    if table_data.get("title"):
        end = get_column_letter(len(headers))
        sheet.merge_cells(start_row=row_start, start_column=1, end_row=row_start, end_column=len(headers))
        title_cell = sheet.cell(row_start, 1, table_data["title"])
        title_cell.font = Font(name="Liberation Sans", size=11, bold=True, color="25445C")
        title_cell.alignment = Alignment(wrap_text=True)
        row_start += 1
    for column, header in enumerate(headers, start=1):
        cell = sheet.cell(row_start, column, header)
        cell.font = Font(name="Liberation Sans", bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="25445C")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for row_number, row in enumerate(table_data["rows"], start=row_start + 1):
        for column, value in enumerate(row, start=1):
            cell = sheet.cell(row_number, column)
            _xlsx_formula_or_value(
                cell,
                value,
                f"{label} row {row_number - row_start}",
                expected_values,
            )
        sheet.row_dimensions[row_number].height = 32
    sheet.row_dimensions[row_start].height = 28
    return row_start + len(table_data["rows"]) + 1


def _model_leaves(value: Any, path: tuple[str, ...] = ()) -> list[tuple[str, str, Any]]:
    result: list[tuple[str, str, Any]] = []
    if isinstance(value, dict):
        if "formula" in value or "expected_value" in value:
            if (
                set(value) - {"formula", "expected_value"}
                or not isinstance(value.get("formula"), str)
                or not value["formula"].startswith("=")
                or "[" in value["formula"]
                or (
                    "expected_value" in value
                    and (
                        isinstance(value["expected_value"], bool)
                        or not isinstance(value["expected_value"], (int, float))
                        or not math.isfinite(value["expected_value"])
                    )
                )
            ):
                _fail(f"Invalid structured formula at {'/'.join(path)}")
            return [("/".join(path[:-1]), path[-1] if path else "Value", value)]
        for key, item in value.items():
            result.extend(_model_leaves(item, path + (str(key),)))
    elif isinstance(value, list):
        for index, item in enumerate(value):
            result.extend(_model_leaves(item, path + (str(index + 1),)))
    else:
        result.append(("/".join(path[:-1]), path[-1] if path else "Value", value))
    return result


def _row_records_table(rows: list[dict[str, Any]], title: str) -> dict[str, Any]:
    headers = list(dict.fromkeys(key for row in rows for key in row))
    if not headers or any(not isinstance(key, str) or not key.strip() for key in headers):
        _fail(f"{title} requires non-empty text keys")
    output_rows = []
    for row in rows:
        output = []
        for header in headers:
            value = row.get(header)
            output.append(
                json.dumps(value, ensure_ascii=False, sort_keys=True)
                if isinstance(value, (dict, list))
                else value
            )
        output_rows.append(output)
    return {"title": title, "headers": headers, "rows": output_rows}


def _write_xlsx(
    path: Path, draft: dict[str, Any], context: dict[str, Any]
) -> tuple[dict[str, float], list[str]]:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Overview"
    _xlsx_title(summary, draft["title"], 4)
    summary.append(["Publication date", draft["publication_date"]])
    summary.append(["Company", context["company"]["data"]["name"]])
    summary.append([])
    summary.append(["Section", "Content"])
    for cell in summary[5]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="25445C")
    for section in draft["sections"]:
        summary.append([section["heading"], section["text"]])
    summary.column_dimensions["A"].width = 34
    summary.column_dimensions["B"].width = 100
    for row in summary.iter_rows(min_row=6):
        summary.row_dimensions[row[0].row].height = min(
            300, max(34, 15 * math.ceil((len(str(row[1].value or "")) + 20) / 110))
        )
    _style_xlsx_sheet(summary)

    formula_expectations: dict[str, float] = {}
    section_sheets: list[str] = []
    used = {"overview"}
    for section in draft["sections"]:
        name = _sheet_name(section["heading"], used)
        section_sheets.append(name)
        sheet = workbook.create_sheet(name)
        _xlsx_title(sheet, draft["title"], 5)
        sheet.cell(2, 1, "Publication date")
        sheet.cell(2, 2, draft["publication_date"])
        sheet.merge_cells("A3:E3")
        sheet.cell(3, 1, section["heading"])
        sheet.cell(3, 1).font = Font(
            name="Liberation Sans", size=13, bold=True, color="25445C"
        )
        sheet.merge_cells("A4:E4")
        text = section["text"]
        if len(text) > 32767:
            _fail(f"{section['heading']}: section text exceeds Excel's cell text limit")
        sheet.cell(4, 1, text)
        sheet.cell(4, 1).alignment = Alignment(wrap_text=True, vertical="top")
        sheet.row_dimensions[4].height = min(
            409, max(60, 15 * math.ceil((len(text) + 45) / 110))
        )
        row_start = 6
        for table_index, table in enumerate(section.get("tables", []), start=1):
            row_start = _xlsx_table(
                sheet,
                table,
                row_start,
                formula_expectations,
                f"{section['heading']} table {table_index}",
            ) + 2
        for column in range(1, max(sheet.max_column, 5) + 1):
            sheet.column_dimensions[get_column_letter(column)].width = 27 if column == 1 else 22
        sheet.column_dimensions["A"].width = 32
        sheet.column_dimensions["B"].width = 28
        sheet.column_dimensions["C"].width = 45
        sheet.print_area = f"A1:{get_column_letter(max(sheet.max_column, 5))}{sheet.max_row}"
        _style_xlsx_sheet(sheet)

    if draft.get("requirement_rows"):
        table_data = _row_records_table(draft["requirement_rows"], "Requirement traceability")
        name = _sheet_name(table_data["title"], used)
        sheet = workbook.create_sheet(name)
        _xlsx_title(sheet, draft["title"], len(table_data["headers"]))
        _xlsx_table(sheet, table_data, 3, formula_expectations, table_data["title"])
        for column, header in enumerate(table_data["headers"], start=1):
            sheet.column_dimensions[get_column_letter(column)].width = min(
                64, max(18, len(header) + 4)
            )
        _style_xlsx_sheet(sheet)

    structured_data = {}
    if draft.get("artifact") is not None:
        structured_data["artifact"] = draft["artifact"]
    if draft.get("model") is not None:
        structured_data["model"] = draft["model"]
    if structured_data:
        leaves = _model_leaves(structured_data)
        if not leaves:
            _fail("Structured workbook data is empty")
        model_table = {
            "title": "Structured source data",
            "headers": ["Source path", "Field", "Value"],
            "rows": [[parent, key, value] for parent, key, value in leaves],
        }
        name = _sheet_name("Structured Data", used)
        sheet = workbook.create_sheet(name)
        _xlsx_title(sheet, "Structured source data", 3)
        _xlsx_table(sheet, model_table, 3, formula_expectations, "Structured source data")
        sheet.column_dimensions["A"].width = 42
        sheet.column_dimensions["B"].width = 34
        sheet.column_dimensions["C"].width = 70
        _style_xlsx_sheet(sheet)

    for sheet in workbook.worksheets:
        sheet.print_area = f"A1:{get_column_letter(max(sheet.max_column, 1))}{sheet.max_row}"
    workbook.properties.title = draft["title"]
    workbook.properties.creator = context["company"]["data"]["name"]
    workbook.properties.subject = draft["publication_date"]
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    workbook.close()
    return formula_expectations, section_sheets


def _pptx_rgb(rgb: tuple[int, int, int]) -> RGBColor:
    return RGBColor(*rgb)


def _shape_box(slide: Any, left: float, top: float, width: float, height: float, fill: tuple[int, int, int]) -> Any:
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(left),
        Inches(top),
        Inches(width),
        Inches(height),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = _pptx_rgb(fill)
    shape.line.fill.background()
    shape.text_frame.clear()
    shape.text_frame.word_wrap = True
    shape.text_frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    shape.text_frame.margin_left = Inches(0.13)
    shape.text_frame.margin_right = Inches(0.13)
    shape.text_frame.margin_top = Inches(0.09)
    shape.text_frame.margin_bottom = Inches(0.07)
    return shape


def _shape_text(
    shape: Any,
    lines: list[str],
    *,
    font_size: int = 18,
    color: tuple[int, int, int] = INK,
    align: Any = PP_ALIGN.LEFT,
    bold: bool = False,
) -> None:
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    for index, line in enumerate(lines):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = line
        paragraph.alignment = align
        paragraph.space_after = Pt(8)
        for run in paragraph.runs:
            run.font.name = "Liberation Sans"
            run.font.size = Pt(font_size)
            run.font.bold = bold
            run.font.color.rgb = _pptx_rgb(color)


def _pptx_add_chart(slide: Any, chart_data: dict[str, Any]) -> None:
    data = CategoryChartData()
    data.categories = chart_data["labels"]
    data.add_series(chart_data.get("title") or "Value", chart_data["values"])
    kind = chart_data.get("kind", "column")
    chart = slide.shapes.add_chart(
        CHART_TYPES[kind],
        Inches(0.8),
        Inches(2.65),
        Inches(11.75),
        Inches(4.25),
        data,
    ).chart
    chart.has_legend = False
    chart.has_title = bool(chart_data.get("title"))
    if chart.has_title:
        chart.chart_title.text_frame.text = chart_data["title"]
    if kind != "pie" and kind != "doughnut":
        chart.value_axis.has_major_gridlines = True
        chart.value_axis.tick_labels.font.name = "Liberation Sans"
        chart.value_axis.tick_labels.font.size = Pt(10)
        chart.category_axis.tick_labels.font.name = "Liberation Sans"
        chart.category_axis.tick_labels.font.size = Pt(10)
    for series in chart.series:
        series.format.fill.solid()
        series.format.fill.fore_color.rgb = _pptx_rgb(BLUE)


def _pptx_add_table(slide: Any, table_data: dict[str, Any]) -> None:
    if len(table_data["headers"]) > 6 or len(table_data["rows"]) > 10:
        _fail(
            "PPTX tables are limited to 6 columns and 10 rows per slide; "
            "use XLSX for larger traceability matrices"
        )
    rows = len(table_data["rows"]) + 1
    columns = len(table_data["headers"])
    shape = slide.shapes.add_table(
        rows, columns, Inches(0.85), Inches(2.35), Inches(11.65), Inches(4.3)
    )
    table = shape.table
    if table_data.get("title"):
        title = slide.shapes.add_textbox(Inches(0.85), Inches(1.95), Inches(11.65), Inches(0.3))
        title.text_frame.text = table_data["title"]
    for row_index, values in enumerate([table_data["headers"], *table_data["rows"]]):
        for column_index, value in enumerate(values):
            cell = table.cell(row_index, column_index)
            cell.text = _simple_value(value, "PPTX table cell")
            cell.fill.solid()
            cell.fill.fore_color.rgb = _pptx_rgb(NAVY if row_index == 0 else WHITE)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            for paragraph in cell.text_frame.paragraphs:
                for run in paragraph.runs:
                    run.font.name = "Liberation Sans"
                    run.font.size = Pt(11 if row_index == 0 else 10)
                    run.font.bold = row_index == 0
                    run.font.color.rgb = _pptx_rgb(WHITE if row_index == 0 else INK)


def _pptx_add_body(slide: Any, body: list[str], layout: str) -> None:
    if layout in {"cards", "grid", "stats", "workflow"}:
        columns = min(3 if layout in {"grid", "stats", "workflow"} else 2, len(body))
        rows = math.ceil(len(body) / columns)
        available_top, available_height = 1.45, 5.45
        gap_x, gap_y = 0.2, 0.18
        width = (11.9 - gap_x * (columns - 1)) / columns
        height = (available_height - gap_y * (rows - 1)) / rows
        for index, text in enumerate(body):
            row, column = divmod(index, columns)
            fill = (PALE, (244, 247, 248), (232, 242, 239))[index % 3]
            shape = _shape_box(
                slide,
                0.72 + column * (width + gap_x),
                available_top + row * (height + gap_y),
                width,
                height,
                fill,
            )
            _shape_text(
                shape,
                [text],
                font_size=22 if layout == "stats" and len(text) < 55 else 16,
                bold=layout == "stats",
            )
    elif layout == "cover":
        bar = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(0.75), Inches(1.55), Inches(0.14), Inches(4.7)
        )
        bar.fill.solid()
        bar.fill.fore_color.rgb = _pptx_rgb(TEAL)
        bar.line.fill.background()
        shape = slide.shapes.add_textbox(Inches(1.15), Inches(1.55), Inches(10.9), Inches(4.8))
        _shape_text(shape, body, font_size=25, color=NAVY, bold=True)
    elif layout in {"statement", "closing"}:
        shape = _shape_box(
            slide,
            0.78,
            1.55,
            11.8,
            5.0,
            NAVY if layout == "closing" else PALE,
        )
        _shape_text(
            shape,
            body,
            font_size=22 if len(body) <= 3 else 17,
            color=WHITE if layout == "closing" else NAVY,
        )
    else:
        shape = slide.shapes.add_textbox(
            Inches(0.85), Inches(1.45), Inches(11.7), Inches(5.3)
        )
        _shape_text(shape, body, font_size=16)


def _write_pptx(
    path: Path, draft: dict[str, Any], context: dict[str, Any]
) -> int:
    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    presentation.core_properties.title = draft["title"]
    presentation.core_properties.author = context["company"]["data"]["name"]
    presentation.core_properties.subject = draft["publication_date"]
    sections = draft["sections"]
    slides = draft.get("slides")
    if slides is None:
        slides = [
            {
                "heading": section["heading"],
                "body": [section["text"]],
                "layout": "content",
            }
            for section in sections
        ]
    blank = presentation.slide_layouts[6]
    for index, (source, section) in enumerate(zip(slides, sections), start=1):
        slide = presentation.slides.add_slide(blank)
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = _pptx_rgb(WHITE)
        if index == 1:
            document_title = slide.shapes.add_textbox(
                Inches(0.72), Inches(0.08), Inches(11.9), Inches(0.24)
            )
            document_title.text_frame.text = draft["title"]
            document_title.text_frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
            for paragraph in document_title.text_frame.paragraphs:
                for run in paragraph.runs:
                    run.font.name = "Liberation Sans"
                    run.font.size = Pt(9)
                    run.font.bold = True
                    run.font.color.rgb = _pptx_rgb(MUTED)
        heading = source.get("heading", source.get("title"))
        title_shape = slide.shapes.add_textbox(
            Inches(0.72), Inches(0.42), Inches(10.6), Inches(0.66)
        )
        title_shape.text_frame.text = heading
        title_shape.text_frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
        for paragraph in title_shape.text_frame.paragraphs:
            for run in paragraph.runs:
                run.font.name = "Liberation Sans"
                run.font.size = Pt(25)
                run.font.bold = True
                run.font.color.rgb = _pptx_rgb(NAVY)

        badge = source.get("badge")
        if badge:
            badge_box = slide.shapes.add_textbox(
                Inches(10.65), Inches(0.48), Inches(2.0), Inches(0.28)
            )
            badge_box.text_frame.text = badge
            for paragraph in badge_box.text_frame.paragraphs:
                paragraph.alignment = PP_ALIGN.RIGHT
                for run in paragraph.runs:
                    run.font.name = "Liberation Sans"
                    run.font.size = Pt(9)
                    run.font.bold = True
                    run.font.color.rgb = _pptx_rgb(TEAL)

        subtitle = source.get("subtitle")
        if subtitle:
            subtitle_box = slide.shapes.add_textbox(
                Inches(0.78), Inches(1.05), Inches(11.7), Inches(0.35)
            )
            subtitle_box.text_frame.text = subtitle
            for paragraph in subtitle_box.text_frame.paragraphs:
                for run in paragraph.runs:
                    run.font.name = "Liberation Sans"
                    run.font.size = Pt(15)
                    run.font.color.rgb = _pptx_rgb(MUTED)

        body = _body_lines(source.get("body"), f"slide {index}")
        layout = source.get("layout", "content")
        if source.get("chart") is not None:
            _pptx_add_chart(slide, source["chart"])
            if body:
                text_box = slide.shapes.add_textbox(
                    Inches(0.85), Inches(1.4), Inches(11.6), Inches(1.05)
                )
                _shape_text(text_box, body, font_size=14)
        else:
            _pptx_add_body(slide, body, layout)

        tables = list(section.get("tables", [])) + list(source.get("tables", []))
        if tables:
            if len(tables) != 1 or source.get("chart") is not None:
                _fail(f"slide {index} has an unsupported table/chart combination")
            _pptx_add_table(slide, tables[0])

        footer = slide.shapes.add_textbox(
            Inches(0.75), Inches(7.08), Inches(11.85), Inches(0.2)
        )
        footer.text_frame.text = (
            f"{context['company']['data']['name']}  |  "
            f"{draft['publication_date']}  |  {index}"
        )
        for paragraph in footer.text_frame.paragraphs:
            paragraph.alignment = PP_ALIGN.RIGHT
            for run in paragraph.runs:
                run.font.name = "Liberation Sans"
                run.font.size = Pt(8)
                run.font.color.rgb = _pptx_rgb(MUTED)
        notes = source.get("notes", section["text"])
        slide.notes_slide.notes_text_frame.text = notes
    path.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(path)
    return len(presentation.slides)


def _check_xlsx(
    path: Path,
    expected_formulas: dict[str, str],
    expected_values: dict[str, float],
) -> dict[str, Any]:
    formulas_book = load_workbook(path, data_only=False)
    cached_book = load_workbook(path, data_only=True)
    formulas: dict[str, str] = {}
    cached: dict[str, Any] = {}
    errors: list[str] = []
    missing: list[str] = []
    try:
        for sheet in formulas_book.worksheets:
            cached_sheet = cached_book[sheet.title]
            for row in sheet.iter_rows():
                for cell in row:
                    if cell.data_type != "f":
                        continue
                    reference = f"{sheet.title}!{cell.coordinate}"
                    formula = str(cell.value)
                    if "[" in formula:
                        _fail(f"External workbook formula is not supported: {reference}")
                    formulas[reference] = formula
                    cached_cell = cached_sheet[cell.coordinate]
                    if cached_cell.value is None:
                        missing.append(reference)
                    else:
                        cached[reference] = cached_cell.value
        for sheet in cached_book.worksheets:
            for row in sheet.iter_rows():
                for cell in row:
                    if cell.data_type == "e" or (
                        isinstance(cell.value, str) and cell.value.upper() in FORMULA_ERRORS
                    ):
                        errors.append(f"{sheet.title}!{cell.coordinate}={cell.value}")
    finally:
        formulas_book.close()
        cached_book.close()
    if formulas != expected_formulas:
        _fail(f"LibreOffice changed XLSX formulas: expected {expected_formulas}, found {formulas}")
    if errors:
        _fail("Formula errors found after LibreOffice recalculation: " + ", ".join(errors))
    if missing:
        _fail("Formula cached values are missing after recalculation: " + ", ".join(missing))
    for reference, expected in expected_values.items():
        actual = cached.get(reference)
        if (
            isinstance(actual, bool)
            or not isinstance(actual, (int, float))
            or not math.isfinite(actual)
            or not math.isclose(actual, expected, rel_tol=1e-9, abs_tol=1e-9)
        ):
            _fail(
                f"Unexpected recalculated result for {reference}: "
                f"expected {expected}, got {actual!r}"
            )
    return {
        "status": "passed",
        "engine": "LibreOffice Calc",
        "formula_count": len(formulas),
        "cached_formula_values": cached,
        "formula_errors": 0,
        "expected_values_checked": len(expected_values),
        "note": (
            "No formulas were present in the supplied draft workbook."
            if not formulas
            else "Formula expressions and cached values were retained and checked."
        ),
    }


class _RunLog:
    def __init__(self, path: Path) -> None:
        if path.exists() or path.is_symlink():
            _fail(f"Refusing to overwrite existing backend log: {path}")
        path.parent.mkdir(parents=True, exist_ok=True)
        path.touch(exist_ok=False)
        self.path = path

    def write(self, text: str) -> None:
        with self.path.open("a", encoding="utf-8", newline="\n") as stream:
            stream.write(text.rstrip() + "\n")


def _soffice_convert(
    source: Path,
    out_dir: Path,
    target: str,
    temporary_root: Path,
    log: _RunLog,
) -> Path:
    executable = shutil.which("soffice")
    if executable is None:
        _fail("LibreOffice soffice is not installed in the backend container")
    if not source.is_file():
        _fail(f"Conversion input does not exist: {source}")
    out_dir.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="lo-profile-", dir=temporary_root) as profile_dir:
        command = [
            executable,
            "--headless",
            "--nologo",
            "--nodefault",
            "--nolockcheck",
            "--norestore",
            f"-env:UserInstallation={Path(profile_dir).as_uri()}",
            "--convert-to",
            target,
            "--outdir",
            str(out_dir),
            str(source),
        ]
        try:
            result = subprocess.run(
                command,
                check=False,
                capture_output=True,
                text=True,
                timeout=300,
            )
        except subprocess.TimeoutExpired as exc:
            _fail(f"LibreOffice timed out converting {source.name} to {target}")
        output = "\n".join(
            part.strip()
            for part in (result.stdout, result.stderr)
            if part.strip()
        )
        log.write("$ " + " ".join(command))
        if output:
            log.write(output)
        if result.returncode != 0:
            _fail(f"LibreOffice failed converting {source.name} to {target}: {output}")
    converted = out_dir / f"{source.stem}.{target.split(':', 1)[0]}"
    if not converted.is_file():
        _fail(f"LibreOffice did not produce {converted}")
    return converted


def _run_soffice_version(log: _RunLog) -> str:
    executable = shutil.which("soffice")
    if executable is None:
        _fail("LibreOffice soffice is not installed in the backend container")
    result = subprocess.run(
        [executable, "--version"],
        check=False,
        capture_output=True,
        text=True,
        timeout=30,
    )
    version = (result.stdout or result.stderr).strip()
    log.write("$ soffice --version")
    log.write(version)
    if result.returncode != 0 or not version:
        _fail("Could not determine the LibreOffice backend version")
    return version


def _copy_new(source: Path, destination: Path) -> None:
    office._ensure_new_file(destination)
    with source.open("rb") as input_stream, destination.open("xb") as output_stream:
        shutil.copyfileobj(input_stream, output_stream)


def _render_and_check(
    artifact: Path,
    document: dict[str, Any],
    draft: dict[str, Any],
    temporary_root: Path,
    log: _RunLog,
) -> dict[str, Any]:
    pdf_dir = temporary_root / "pdf" / document["id"]
    pdf = _soffice_convert(artifact, pdf_dir, "pdf", temporary_root, log)
    expected_text = [document["title"]]
    expected_text.extend(section["heading"] for section in draft["sections"])
    pdf_validation = office.validate_pdf(pdf, tuple(expected_text))
    if document["format"] == "pptx" and pdf_validation["pages"] != len(
        draft.get("slides", draft["sections"])
    ):
        _fail(
            f"{document['id']} rendered {pdf_validation['pages']} PDF pages for "
            f"{len(draft.get('slides', draft['sections']))} slides"
        )
    previews = office._rasterize_pdf(
        pdf,
        temporary_root / "previews" / document["id"],
        document["id"],
        pdf_validation["pages"],
    )
    return {
        "pdf": pdf,
        "pdf_validation": pdf_validation,
        "previews": previews,
    }


def _document_plan_and_inputs(
    contract: Any,
    canon_path: Path,
    plan_path: Path,
    contexts_path: Path,
    drafts_path: Path,
    document_ids: list[str] | None,
) -> tuple[dict[str, Any], dict[str, Any], list[dict[str, Any]], dict[str, Any], dict[str, Any]]:
    canon = _json_file(canon_path)
    plan = _json_file(plan_path)
    contract.validate_canon(canon)
    documents = contract.validate_plan(canon, plan)
    selected = list(documents) if not document_ids else document_ids
    if not selected or len(selected) != len(set(selected)):
        _fail("Select one or more distinct document IDs")
    missing = set(selected) - documents.keys()
    if missing:
        _fail(f"Document IDs are not in the plan: {sorted(missing)}")
    context_files = _json_paths(contexts_path, "Context", set(documents))
    draft_files = _json_paths(drafts_path, "Draft", set(documents))
    prepared = []
    for identifier in selected:
        if identifier not in context_files:
            _fail(f"Missing date-filtered context for {identifier}")
        if identifier not in draft_files:
            _fail(f"Missing grounded draft for {identifier}")
        document = documents[identifier]
        context = _json_file(context_files[identifier])
        draft = _json_file(draft_files[identifier])
        expected_context = contract.get_context(canon, plan, identifier)
        if contract.fingerprint(context) != contract.fingerprint(expected_context):
            _fail(f"{identifier} context does not match its date-filtered canon/plan")
        if draft.get("document_id") != identifier:
            _fail(f"{identifier} draft filename and document_id do not match")
        contract.check_draft(canon, plan, draft)
        if draft.get("document_id") != identifier:
            _fail(f"{identifier} draft refers to another plan item")
        _validate_draft_structure(draft, document)
        contract.validate_content_voice(draft, identifier)
        prepared.append(
            {
                "document": document,
                "context": context,
                "draft": draft,
                "context_path": context_files[identifier],
                "draft_path": draft_files[identifier],
            }
        )
    return canon, plan, prepared, context_files, draft_files


def _prepare_corpus_root(
    root: Path,
    canon_path: Path,
    plan_path: Path,
    prepared: list[dict[str, Any]],
    contract: Any,
) -> tuple[Path, Path, Path]:
    _reject_link(root, "Output corpus root")
    root.mkdir(parents=True, exist_ok=True)
    resolved_root = root.resolve()
    existing = {entry.name for entry in root.iterdir()}
    extra = existing - ROOT_ENTRIES
    if extra:
        _fail(f"Output corpus root contains unsupported top-level entries: {sorted(extra)}")
    docs_dir = root / "documents"
    build_dir = root / "build"
    _ensure_directory(docs_dir, "Output documents directory")
    _ensure_directory(build_dir, "Output build directory")

    _stage_file(canon_path, root / "canon.json", contract)
    _stage_file(plan_path, build_dir / "plan.json", contract)
    request_path = build_dir / "request.json"
    if request_path.exists():
        if contract.fingerprint(_json_file(request_path)) != contract.fingerprint(
            _json_file(root / "canon.json")["request"]
        ):
            _fail(f"Output request.json does not match canon.json: {request_path}")
    else:
        _write_json_new(request_path, _json_file(root / "canon.json")["request"])

    contexts_dir = build_dir / "contexts"
    drafts_dir = build_dir / "drafts"
    _ensure_directory(contexts_dir, "Output contexts directory")
    _ensure_directory(drafts_dir, "Output drafts directory")
    for item in prepared:
        _stage_file(item["context_path"], contexts_dir / f"{item['document']['id']}.json", contract)
        _stage_file(item["draft_path"], drafts_dir / f"{item['document']['id']}.json", contract)
    contract.validate_output_layout(_json_file(root / "canon.json"), root)
    return root, docs_dir, build_dir


def _preflight_output_targets(
    docs_dir: Path,
    build_dir: Path,
    prepared: list[dict[str, Any]],
    is_complete_plan: bool,
) -> tuple[Path, Path, Path]:
    existing = list(docs_dir.rglob("*"))
    for path in existing:
        if path.is_symlink():
            _fail(f"Output documents directory contains a symbolic link: {path}")
        if path.is_file():
            _fail(f"Refusing to overwrite an existing document: {path}")
    renders = build_dir / "renders"
    _reject_link(renders, "Output render directory")
    if renders.exists() and any(renders.iterdir()):
        _fail(f"Refusing to overwrite existing render artifacts: {renders}")
    if renders.exists() and not renders.is_dir():
        _fail(f"Output render path is not a directory: {renders}")
    manifest = build_dir / (
        "manifest.json" if is_complete_plan else "backend-pilot-manifest.json"
    )
    evidence = build_dir / (
        "backend-evidence.json" if is_complete_plan else "backend-pilot-evidence.json"
    )
    log_path = build_dir / "backend-export.log"
    for path in (manifest, evidence, log_path):
        if path.exists() or path.is_symlink():
            _fail(f"Refusing to overwrite existing backend output: {path}")
    return renders, manifest, evidence


def export_documents(
    canon_path: Path,
    plan_path: Path,
    contexts_path: Path,
    drafts_path: Path,
    root: Path,
    document_ids: list[str] | None = None,
) -> dict[str, Any]:
    for path, label in (
        (canon_path, "Canon"),
        (plan_path, "Plan"),
        (contexts_path, "Contexts"),
        (drafts_path, "Drafts"),
    ):
        _reject_link(path, label)
    if not canon_path.is_file() or not plan_path.is_file():
        _fail("Canon and plan must be regular input files")
    if not contexts_path.is_dir() or not drafts_path.is_dir():
        _fail("Contexts and drafts must be existing directories")
    if Path(root).is_symlink():
        _fail(f"Output corpus root must not be a symbolic link: {root}")

    contract = office._company_contract()
    canon, plan, prepared, _, _ = _document_plan_and_inputs(
        contract,
        canon_path,
        plan_path,
        contexts_path,
        drafts_path,
        document_ids,
    )
    documents = contract.validate_plan(canon, plan)
    complete_plan = {item["document"]["id"] for item in prepared} == set(documents)
    root, docs_dir, build_dir = _prepare_corpus_root(
        Path(root), canon_path, plan_path, prepared, contract
    )
    contract.require_build_path(build_dir / "drafts", root)
    renders, manifest_path, evidence_path = _preflight_output_targets(
        docs_dir, build_dir, prepared, complete_plan
    )
    log = _RunLog(build_dir / "backend-export.log")
    version = _run_soffice_version(log)
    initial_input_hashes = {
        item["document"]["id"]: {
            "context_sha256": _sha256(item["context_path"]),
            "draft_sha256": _sha256(item["draft_path"]),
        }
        for item in prepared
    }

    generated: list[dict[str, Any]] = []
    with tempfile.TemporaryDirectory(prefix=".backend-stage-", dir=build_dir) as temp_name:
        temporary_root = Path(temp_name)
        generated_docs = temporary_root / "documents"
        generated_docs.mkdir()
        for item in prepared:
            document = item["document"]
            draft = item["draft"]
            context = item["context"]
            temp_artifact = generated_docs / document["filename"]
            formulas: dict[str, str] = {}
            expected_values: dict[str, float] = {}
            slides = None
            if document["format"] == "docx":
                _write_docx(temp_artifact, draft, context)
            elif document["format"] == "xlsx":
                expected_values, _ = _write_xlsx(temp_artifact, draft, context)
                formulas_before = office._read_formulas(temp_artifact)
                recalc_dir = temporary_root / "recalculated"
                recalculated = _soffice_convert(
                    temp_artifact,
                    recalc_dir,
                    "xlsx:Calc MS Excel 2007 XML",
                    temporary_root,
                    log,
                )
                formulas = formulas_before
                recalc_result = _check_xlsx(recalculated, formulas_before, expected_values)
                temp_artifact.unlink()
                shutil.copyfile(recalculated, temp_artifact)
            elif document["format"] == "pptx":
                slides = _write_pptx(temp_artifact, draft, context)
            else:
                _fail(f"{document['id']} has unsupported output format {document['format']!r}")

            contract.audit_package(temp_artifact, document["format"])
            rendered = _render_and_check(
                temp_artifact, document, draft, temporary_root, log
            )
            if document["format"] == "xlsx":
                recalc_result = _check_xlsx(temp_artifact, formulas, expected_values)
            generated.append(
                {
                    "item": item,
                    "temporary_artifact": temp_artifact,
                    "temporary_pdf": rendered["pdf"],
                    "temporary_previews": rendered["previews"],
                    "pdf_validation": rendered["pdf_validation"],
                    "recalculation": recalc_result if document["format"] == "xlsx" else None,
                    "slide_count": slides,
                }
            )

        for result in generated:
            item = result["item"]
            document = item["document"]
            final_artifact = docs_dir / document["filename"]
            _copy_new(result["temporary_artifact"], final_artifact)
            render_dir = renders / document["id"]
            _ensure_directory(render_dir, "Output per-document render directory")
            final_pdf = render_dir / "document.pdf"
            _copy_new(result["temporary_pdf"], final_pdf)
            final_previews = []
            for preview in result["temporary_previews"]:
                final_preview = render_dir / preview.name
                _copy_new(preview, final_preview)
                final_previews.append(final_preview)
            result["artifact"] = final_artifact
            result["pdf"] = final_pdf
            result["previews"] = final_previews

    manifest = {
        "schema_version": 1,
        "canon_sha256": contract.fingerprint(canon),
        "plan_sha256": contract.fingerprint(plan),
        "scope": "complete plan" if complete_plan else "selected plan documents only",
        "documents": [
            {
                "document_id": result["item"]["document"]["id"],
                "path": result["artifact"].relative_to(root).as_posix(),
                "sha256": _sha256(result["artifact"]),
            }
            for result in generated
        ],
    }
    evidence = {
        "schema_version": 1,
        "scope": "complete plan" if complete_plan else "selected plan documents only; pilot evidence",
        "backend": {
            "name": "LibreOffice",
            "version": version,
            "containerized": True,
            "office_fidelity_claim": False,
        },
        "canon_sha256": contract.fingerprint(canon),
        "plan_sha256": contract.fingerprint(plan),
        "documents": [],
        "reviews": {
            "contract_and_package_integrity": "passed",
            "render_and_text_checks": "passed",
            "visual_review": "required; page previews were generated but not independently reviewed",
            "semantic_review": "required",
            "genre_review": "required",
            "depth_review": "required",
            "office_compatibility": "not assessed",
        },
    }
    for result in generated:
        item = result["item"]
        document = item["document"]
        evidence["documents"].append(
            {
                "document_id": document["id"],
                "format": document["format"],
                "title": document["title"],
                "context_sha256": initial_input_hashes[document["id"]]["context_sha256"],
                "draft_sha256": initial_input_hashes[document["id"]]["draft_sha256"],
                "artifact": result["artifact"].relative_to(root).as_posix(),
                "artifact_sha256": _sha256(result["artifact"]),
                "pdf": result["pdf"].relative_to(root).as_posix(),
                "pdf_validation": result["pdf_validation"],
                "previews": [path.relative_to(root).as_posix() for path in result["previews"]],
                "slide_count": result["slide_count"],
                "recalculation": result["recalculation"],
            }
        )
    _write_json_new(manifest_path, manifest)
    _write_json_new(evidence_path, evidence)
    result = {
        "status": "rendered; editorial and visual reviews remain required",
        "output_root": str(root),
        "manifest": str(manifest_path),
        "evidence": str(evidence_path),
        "log": str(log.path),
        "documents": evidence["documents"],
    }
    print(json.dumps(result, indent=2, ensure_ascii=False))
    return result


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description=(
            "Validate prepared company canon/plan/context/draft artifacts and export "
            "the selected DOCX, XLSX, and PPTX documents in this container."
        )
    )
    parser.add_argument("--canon", type=Path, required=True, help="Validated corpus canon.json")
    parser.add_argument("--plan", type=Path, required=True, help="Frozen build/plan.json")
    parser.add_argument("--contexts", type=Path, required=True, help="Date-filtered context directory")
    parser.add_argument("--drafts", type=Path, required=True, help="Grounded draft directory")
    parser.add_argument(
        "--root",
        "--output-root",
        dest="root",
        type=Path,
        required=True,
        help="Output corpus root; inputs are staged without overwriting existing files",
    )
    parser.add_argument(
        "--document",
        action="append",
        dest="document_ids",
        help="Optional planned document ID to export; repeat for a bounded pilot",
    )
    args = parser.parse_args(argv)
    try:
        export_documents(
            args.canon,
            args.plan,
            args.contexts,
            args.drafts,
            args.root,
            args.document_ids,
        )
        return 0
    except Exception as exc:
        if isinstance(exc, (KeyboardInterrupt, SystemExit)):
            raise
        print(f"error: {exc}", file=sys.stderr)
        return 1
