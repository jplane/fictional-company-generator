import json
import os
import pwd
import shutil
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn

import office_backend
from openpyxl import Workbook, load_workbook

from office_backend import (
    BackendError,
    _company_contract,
    inspect_recalculated_workbook,
    recalculate_xlsx,
    run_demo,
    validate_pdf,
)
from document_export import _write_docx, export_documents


class LibreOfficeBackendSmokeTests(unittest.TestCase):
    def test_container_runs_as_non_root(self):
        self.assertGreater(os.geteuid(), 0)
        self.assertEqual(pwd.getpwuid(os.geteuid()).pw_shell, "/bin/bash")

    def test_contract_helper_mount_is_read_only(self):
        helper = Path("/app/company_contract.py")
        self.assertTrue(helper.is_file())
        with self.assertRaises(OSError):
            with helper.open("r+b"):
                pass

    def test_contract_helper_loads_with_workspace_mounted_source(self):
        with patch(
            "office_backend.__file__",
            "/workspace/prototypes/libreoffice-backend/office_backend.py",
        ):
            contract = _company_contract()
        self.assertEqual(Path(contract.__file__), Path("/app/company_contract.py"))

    def test_root_execution_is_rejected(self):
        with patch("office_backend.os.geteuid", return_value=0):
            with self.assertRaisesRegex(BackendError, "refuses to run as root"):
                run_demo(Path("/tmp/root-run-must-not-create"))

    def test_date_filtered_docx_xlsx_contract_pilot(self):
        with tempfile.TemporaryDirectory() as temporary_dir:
            output = Path(temporary_dir) / "pilot"
            results = run_demo(output)

            canon = json.loads((output / "canon.json").read_text(encoding="utf-8"))
            plan = json.loads(
                (output / "build" / "plan.json").read_text(encoding="utf-8")
            )
            self.assertEqual(
                {document["format"] for document in plan["documents"]},
                {"docx", "xlsx"},
            )
            self.assertEqual(len(plan["documents"]), 2)
            for document in plan["documents"]:
                context = json.loads(
                    (
                        output
                        / "build"
                        / "contexts"
                        / f"{document['id']}.json"
                    ).read_text(encoding="utf-8")
                )
                draft = json.loads(
                    (
                        output
                        / "build"
                        / "drafts"
                        / f"{document['id']}.json"
                    ).read_text(encoding="utf-8")
                )
                self.assertEqual(context["document"], document)
                self.assertEqual(draft["document_id"], document["id"])
                self.assertEqual(draft["canon_sha256"], plan["canon_sha256"])
                visible_ids = {record["id"] for record in context["records"]}
                anchor = next(
                    record
                    for record in canon["records"]
                    if record["id"] == document["anchor_ids"][0]
                )
                future_launches = [
                    record
                    for record in canon["records"]
                    if record["kind"] == "milestone"
                    and record["data"]["product_id"]
                    == anchor["data"]["product_id"]
                    and record["data"]["status"] == "launched"
                    and record["known_on"] > document["publication_date"]
                ]
                self.assertTrue(future_launches)
                self.assertTrue(
                    all(record["id"] not in visible_ids for record in future_launches)
                )

            for document_result in results["documents"]:
                self.assertTrue((output / document_result["artifact"]).is_file())
                self.assertGreater(document_result["pdf_validation"]["pages"], 0)
                for image in document_result["visual_pages"]:
                    self.assertTrue((output / image).read_bytes().startswith(b"\x89PNG"))
            self.assertEqual(
                results["contract_validation"]["checked_drafts"], 2
            )
            self.assertTrue(results["contract_validation"]["office_package_integrity"])
            self.assertEqual(results["office_compatibility_qa"], "not performed")

            workbook_path = output / "documents" / next(
                item["filename"] for item in plan["documents"] if item["format"] == "xlsx"
            )
            recalculated = output / "build" / "renders" / (
                "recalculated-" + next(
                    item["filename"]
                    for item in plan["documents"]
                    if item["format"] == "xlsx"
                )
            )
            formulas_book = load_workbook(workbook_path, data_only=False)
            cached_book = load_workbook(recalculated, data_only=True)
            formulas = {
                f"{sheet.title}!{cell.coordinate}": cell.value
                for sheet in formulas_book.worksheets
                for row in sheet.iter_rows()
                for cell in row
                if cell.data_type == "f"
            }
            cached = {
                f"{sheet.title}!{cell.coordinate}": cell.value
                for sheet in cached_book.worksheets
                for row in sheet.iter_rows()
                for cell in row
                if cell.data_type == "f" or cell.value is not None
            }
            formulas_book.close()
            cached_book.close()
            self.assertEqual(
                formulas,
                {
                    "Validation!B3": "=COUNTA(Requirements!A3:A4)",
                    "Validation!B4": "=COUNTA('Acceptance Criteria'!B3:B6)",
                },
            )
            self.assertEqual(cached["Validation!B3"], 2)
            self.assertEqual(cached["Validation!B4"], 4)

    def test_formula_errors_are_reported_after_recalculation(self):
        with tempfile.TemporaryDirectory() as temporary_dir:
            root = Path(temporary_dir)
            source = root / "broken.xlsx"
            recalculated = root / "recalculated" / "broken.xlsx"
            workbook = Workbook()
            workbook.active["A1"] = "=1/0"
            workbook.save(source)
            workbook.close()

            recalculate_xlsx(source, recalculated)
            with self.assertRaisesRegex(BackendError, "Formula errors found"):
                inspect_recalculated_workbook(recalculated)

    def test_existing_output_directory_is_never_overwritten(self):
        with tempfile.TemporaryDirectory() as temporary_dir:
            output = Path(temporary_dir) / "existing"
            output.mkdir()
            sentinel = output / "keep.txt"
            sentinel.write_text("preserve me", encoding="utf-8")

            with self.assertRaisesRegex(BackendError, "non-empty output directory"):
                run_demo(output)

            self.assertEqual(sentinel.read_text(encoding="utf-8"), "preserve me")
            self.assertEqual(list(output.iterdir()), [sentinel])

    def test_recalculation_refuses_to_overwrite_destination(self):
        with tempfile.TemporaryDirectory() as temporary_dir:
            root = Path(temporary_dir)
            source = root / "source.xlsx"
            destination = root / "existing.xlsx"
            workbook = Workbook()
            workbook.active["A1"] = "=1+1"
            workbook.save(source)
            workbook.close()
            destination.write_bytes(b"preserve me")

            with self.assertRaisesRegex(BackendError, "Refusing to overwrite"):
                recalculate_xlsx(source, destination)

            self.assertEqual(destination.read_bytes(), b"preserve me")

    def test_malformed_pdf_is_rejected(self):
        with tempfile.TemporaryDirectory() as temporary_dir:
            path = Path(temporary_dir) / "invalid.pdf"
            path.write_bytes(b"%PDF-1.4\nnot a complete PDF")

            with self.assertRaisesRegex(BackendError, "header or end marker"):
                validate_pdf(path, ("Required text",))

    def test_cli_requires_prepared_artifact_paths(self):
        with self.assertRaises(SystemExit) as raised:
            office_backend.main([])
        self.assertEqual(raised.exception.code, 2)


class ArtifactDrivenExportTests(unittest.TestCase):
    def setUp(self):
        self.temporary_dir = tempfile.TemporaryDirectory()
        self.root = Path(self.temporary_dir.name)
        self.inputs = self.root / "inputs"
        self.build = self.inputs / "build"
        self.contexts = self.build / "contexts"
        self.drafts = self.build / "drafts"
        self.documents_dir = self.root / "output"
        self.contexts.mkdir(parents=True)
        self.drafts.mkdir()
        self.contract = _company_contract()
        request = {
            "company_name": "Lumenvale Systems",
            "industry": "industrial analytics software",
            "start_year": 2022,
            "years": 1,
            "documents": 3,
            "seed": 42,
            "content_depth": "light",
            "document_types": ["annual_report"],
            "formats": ["docx", "xlsx", "pptx"],
        }
        self.canon = self.contract.seed_canon(request)
        self.plan = self.contract.create_plan(self.canon)
        self.contract.save(self.inputs / "canon.json", self.canon)
        self.contract.save(self.build / "plan.json", self.plan)
        self.contexts_by_id = {}
        self.drafts_by_id = {}
        for item in self.plan["documents"]:
            context = self.contract.get_context(self.canon, self.plan, item["id"])
            anchor = next(
                record
                for record in context["records"]
                if record["id"] == item["anchor_ids"][0]
            )
            claim = {
                "id": "claim-anchor",
                "record_id": anchor["id"],
                "path": ["data"],
                "value": anchor["data"],
            }
            result_text = (
                f"Recognized revenue for the recorded period was "
                f"${anchor['data']['revenue_usd']:,}."
            )
            bodies = {
                "financial_results": result_text,
                "business_review": "The review uses the approved annual operating record.",
                "organization": "The date-filtered context controls the applicable organization.",
                "risks": "The annual record is the basis for the scoped review.",
            }
            sections = []
            for role in item["required_sections"]:
                section = {
                    "role": role,
                    "heading": role.replace("_", " ").title(),
                    "text": bodies[role],
                    "claim_ids": ["claim-anchor"],
                }
                if role == "financial_results" and item["format"] in {"docx", "xlsx"}:
                    row = ["Recognized revenue", f"${anchor['data']['revenue_usd']:,}"]
                    if item["format"] == "xlsx":
                        row = [
                            "Reconciliation check",
                            {"formula": "=SUM(2, 3)", "expected_value": 5},
                        ]
                    section["tables"] = [
                        {
                            "title": "Annual source figures",
                            "headers": ["Measure", "Actual"],
                            "rows": [row],
                        }
                    ]
                sections.append(section)
            draft = {
                "document_id": item["id"],
                "canon_sha256": context["canon_sha256"],
                "plan_sha256": context["plan_sha256"],
                "title": item["title"],
                "publication_date": item["publication_date"],
                "content_depth": item["content_depth"],
                "claims": [claim],
                "sections": sections,
            }
            if item["format"] == "pptx":
                draft["slides"] = [
                    {
                        "heading": section["heading"],
                        "body": [section["text"]],
                        "layout": "chart" if index == 0 else "statement",
                        **(
                            {
                                "chart": {
                                    "labels": ["Revenue", "COGS", "Operating costs"],
                                    "values": [
                                        anchor["data"]["revenue_usd"],
                                        anchor["data"]["cogs_usd"],
                                        anchor["data"]["opex_usd"],
                                    ],
                                },
                                "notes": section["text"],
                            }
                            if index == 0
                            else {}
                        ),
                    }
                    for index, section in enumerate(sections)
                ]
            self.contract.check_draft(self.canon, self.plan, draft)
            self.contract.save(
                self.contexts / f"{item['id']}.json",
                context,
            )
            self.contract.save(self.drafts / f"{item['id']}.json", draft)
            self.contexts_by_id[item["id"]] = context
            self.drafts_by_id[item["id"]] = draft

    def tearDown(self):
        self.temporary_dir.cleanup()

    def _export(self, selected=None, root=None):
        return export_documents(
            self.inputs / "canon.json",
            self.build / "plan.json",
            self.contexts,
            self.drafts,
            root or self.documents_dir,
            selected,
        )

    def test_exports_all_prepared_formats_with_render_and_recalc_evidence(self):
        result = self._export()
        self.assertEqual(len(result["documents"]), 3)
        self.assertTrue((self.documents_dir / "build" / "manifest.json").is_file())
        evidence = json.loads(
            (self.documents_dir / "build" / "backend-evidence.json").read_text(
                encoding="utf-8"
            )
        )
        self.assertEqual(
            {entry["format"] for entry in evidence["documents"]},
            {"docx", "xlsx", "pptx"},
        )
        for entry in evidence["documents"]:
            artifact = self.documents_dir / entry["artifact"]
            pdf = self.documents_dir / entry["pdf"]
            self.assertTrue(artifact.is_file())
            self.assertTrue(pdf.is_file())
            self.assertGreater(entry["pdf_validation"]["pages"], 0)
            self.assertTrue(entry["previews"])
            self.assertTrue(
                all(
                    (self.documents_dir / preview).read_bytes().startswith(
                        b"\x89PNG\r\n\x1a\n"
                    )
                    for preview in entry["previews"]
                )
            )
            self.assertEqual(entry["artifact_sha256"], office_backend._sha256(artifact))
        xlsx = next(item for item in evidence["documents"] if item["format"] == "xlsx")
        self.assertEqual(xlsx["recalculation"]["formula_count"], 1)
        self.assertEqual(xlsx["recalculation"]["cached_formula_values"], {
            "Financial Results!B8": 5
        })
        pptx = next(item for item in evidence["documents"] if item["format"] == "pptx")
        self.assertEqual(pptx["slide_count"], 4)
        self.assertEqual(pptx["pdf_validation"]["pages"], 4)
        self.assertEqual(evidence["reviews"]["visual_review"].split(";")[0], "required")

    def test_docx_wide_tables_use_landscape_and_repeat_headers(self):
        docx_id = next(
            item["id"] for item in self.plan["documents"] if item["format"] == "docx"
        )
        draft = json.loads(json.dumps(self.drafts_by_id[docx_id]))
        draft["sections"][0]["tables"] = [
            {
                "title": "Quarterly results",
                "headers": [
                    "Quarter",
                    "Revenue",
                    "COGS",
                    "Operating expenses",
                    "Operating income",
                    "Sequential change",
                ],
                "rows": [["Q1", "$10,000", "$3,000", "$4,000", "$3,000", "-"]],
            }
        ]
        path = self.root / "wide-table.docx"
        _write_docx(path, draft, self.contexts_by_id[docx_id])

        exported = Document(path)
        self.assertEqual(
            [section.orientation for section in exported.sections],
            [WD_ORIENT.PORTRAIT, WD_ORIENT.LANDSCAPE, WD_ORIENT.PORTRAIT],
        )
        self.assertIsNotNone(
            exported.tables[0]._tbl.tr_lst[0].trPr.find(qn("w:tblHeader"))
        )
        self.assertIsNotNone(
            exported.tables[0]._tbl.tr_lst[1].trPr.find(qn("w:cantSplit"))
        )
        expected_footer = (
            f"{self.contexts_by_id[docx_id]['company']['data']['name']} "
            f"| {draft['publication_date']}"
        )
        self.assertEqual(
            [section.footer.paragraphs[0].text for section in exported.sections],
            [expected_footer] * 3,
        )

    def test_selected_plan_items_get_explicit_pilot_manifest(self):
        selected = ["doc-0003", "doc-0001"]
        root = self.root / "selected"
        result = self._export(selected, root)
        self.assertEqual(len(result["documents"]), 2)
        self.assertTrue((root / "build" / "backend-pilot-manifest.json").is_file())
        self.assertFalse((root / "build" / "manifest.json").exists())
        self.assertEqual(
            len(list((root / "documents").glob("*"))),
            2,
        )
        evidence = json.loads(
            (root / "build" / "backend-pilot-evidence.json").read_text(encoding="utf-8")
        )
        self.assertEqual(evidence["scope"], "selected plan documents only; pilot evidence")

    def test_context_mismatch_is_rejected_before_output_creation(self):
        context_path = self.contexts / "doc-0001.json"
        original = json.loads(context_path.read_text(encoding="utf-8"))
        original["document"]["title"] = "Changed title"
        context_path.write_text(json.dumps(original), encoding="utf-8")
        root = self.root / "bad-context"
        with self.assertRaisesRegex(BackendError, "context does not match"):
            self._export(["doc-0001"], root)
        self.assertFalse(root.exists())

    def test_unsupported_format_structure_is_reported_not_dropped(self):
        item = self.plan["documents"][0]
        draft_path = self.drafts / f"{item['id']}.json"
        draft = json.loads(draft_path.read_text(encoding="utf-8"))
        draft["sections"][0]["unknown_chart"] = {"values": [1]}
        draft_path.write_text(json.dumps(draft), encoding="utf-8")
        with self.assertRaisesRegex(BackendError, "unsupported structures"):
            self._export([item["id"]], self.root / "unsupported")

    def test_existing_artifacts_are_never_overwritten(self):
        self._export()
        artifact = next((self.documents_dir / "documents").iterdir())
        before = artifact.read_bytes()
        with self.assertRaisesRegex(BackendError, "Refusing to overwrite an existing document"):
            self._export()
        self.assertEqual(artifact.read_bytes(), before)


if __name__ == "__main__":
    unittest.main()
